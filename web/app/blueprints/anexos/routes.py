# web/app/blueprints/anexos/routes.py

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from flask import current_app, flash, redirect, render_template, request, url_for, abort
from flask_login import current_user, login_required

from ...auth.decorators import role_required
from ...config import generar_nombre_documento
from ...extensions import db
from ...models import (
    AnexoExtensionContrato,
    Contrato,
    EventoLaboral,  # legacy (si aún existe para otras vistas)
)
from ...services.docx_engine import replace_placeholders
from ...services.paths_nextcloud import get_nc_paths
from ...services.pdf_convert import convert_docx_to_pdf
from ...services.storage_nextcloud_fs import ensure_dir
from ...templates.utils.dates import fecha_larga_es
from ...utils import parse_date
from . import bp

from docx import Document


# ==========================
# Helpers: acceso por obra (alineado a tu lógica global)
# ==========================
def _require_contrato_access(contrato: Contrato) -> None:
    """
    Control de acceso por obra asociada al contrato.
    - ADMIN: acceso total
    - Otros: solo si la obra del contrato está asignada al usuario
    """
    if current_user.has_role("ADMIN"):
        return

    obra_id = getattr(contrato, "obra_id", None)
    if obra_id is None:
        abort(403)

    if not current_user.can_access_obra(int(obra_id)):
        abort(403)


# ==========================
# DOCX builder: Anexo Extensión
# ==========================
def build_anexo_extension_docx(template_path: Path, variables: dict[str, str], output_path: Path) -> Path:
    doc = Document(str(template_path))
    replace_placeholders(doc, variables)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _get_tpl_anexo_extension() -> Path:
    templates_root = Path(current_app.root_path) / "document_templates" / "anexos"
    tpl = templates_root / "anexo_extension_contrato.docx"
    return tpl


def _build_variables_anexo_extension(
    contrato: Contrato,
    fecha_anexo: date,
    anterior: date | None,
    nueva_fecha_termino: date,
    observaciones: str | None,
) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    rut_digits = (getattr(trabajador, "rut", "") or "").strip()
    dv = (getattr(trabajador, "dv", "") or "").strip().upper()
    rut_completo = f"{rut_digits}-{dv}" if rut_digits and dv else rut_digits

    variables: dict[str, str] = {
        "{{OBRA_NOMBRE}}": (obra.nombre if obra else ""),
        "{{OBRA_COMUNA}}": (obra.comuna if obra else ""),
        "{{ANEXO_FECHA}}": fecha_larga_es(fecha_anexo),

        "{{EMPLEADOR_RAZON_SOCIAL}}": (empleador.razon_social or ""),
        "{{EMPLEADOR_RUT}}": (getattr(empleador, "rut", "") or ""),
        "{{REP_LEGAL_NOMBRE}}": (getattr(empleador, "nombre_rep_legal", "") or ""),
        "{{REP_LEGAL_RUT}}": (getattr(empleador, "rut_rep_legal", "") or ""),
        "{{EMPLEADOR_DIRECCION}}": (empleador.direccion or ""),
        "{{EMPLEADOR_COMUNA}}": (empleador.comuna or ""),

        "{{TRABAJADOR_NOMBRES}}": (trabajador.nombres or ""),
        "{{TRABAJADOR_AP_PATERNO}}": (trabajador.ap_paterno or ""),
        "{{TRABAJADOR_AP_MATERNO}}": (trabajador.ap_materno or ""),
        "{{TRABAJADOR_RUT}}": rut_completo,

        "{{CONTRATO_FECHA_INICIO}}": fecha_larga_es(contrato.fecha_inicio) if contrato.fecha_inicio else "",
        "{{CONTRATO_FECHA_TERMINO_ANTERIOR}}": fecha_larga_es(anterior) if anterior else "",
        "{{CONTRATO_FECHA_TERMINO_NUEVA}}": fecha_larga_es(nueva_fecha_termino),
        "{{CARGO_NOMBRE}}": (contrato.cargo.nombre if contrato.cargo else ""),

        "{{ANEXO_OBSERVACIONES}}": (observaciones or ""),
    }
    return variables


def _generar_archivos_anexo_extension_nextcloud(
    anexo: AnexoExtensionContrato,
    contrato: Contrato,
    formato: str,
) -> None:
    """
    Genera/Regenera archivos (DOCX/PDF) en Nextcloud para un anexo existente.
    Importante: NO crea registros; solo actualiza docx/pdf_* del mismo anexo.
    """
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")
    if not getattr(contrato, "obra", None):
        raise RuntimeError("El contrato no tiene obra asignada (no se puede determinar centro de costo).")

    trabajador = contrato.trabajador
    obra = contrato.obra

    centro_costo = (getattr(obra, "centro_costo", "") or "").strip()
    if not centro_costo:
        raise RuntimeError(
            f"La obra '{getattr(obra, 'nombre', 'SIN_NOMBRE')}' no tiene centro de costo definido."
        )

    tpl = _get_tpl_anexo_extension()
    if not tpl.exists():
        raise RuntimeError("No se encontró la plantilla anexo_extension_contrato.docx en document_templates/anexos.")

    # Variables
    anterior = anexo.fecha_termino_anterior
    nueva_fecha_termino = anexo.fecha_termino_nueva
    if not nueva_fecha_termino:
        raise RuntimeError("El anexo no tiene fecha_termino_nueva definida.")

    variables = _build_variables_anexo_extension(
        contrato=contrato,
        fecha_anexo=anexo.fecha_anexo or date.today(),
        anterior=anterior,
        nueva_fecha_termino=nueva_fecha_termino,
        observaciones=anexo.observaciones,
    )

    # Destino Nextcloud (RUTA LEGACY por CENTRO DE COSTO)
    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_extension_{contrato.id}_{anexo.id}_{stamp}.docx"
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)

    ensure_dir(dest_dir)

    # Construye docx
    build_anexo_extension_docx(tpl, variables, docx_tmp)

    # Nombres finales (estables por fecha_anexo; regenerar sobreescribe)
    fecha_ref = anexo.fecha_anexo or date.today()

    nombre_docx = generar_nombre_documento(
        tipo="ANEXO_EXTENSION_CONTRATO",
        ap_paterno=trabajador.ap_paterno or "SIN_APELLIDO",
        fecha_ref=fecha_ref,
        extension="docx",
    )
    docx_final = dest_dir / nombre_docx
    docx_final.write_bytes(docx_tmp.read_bytes())

    # PDF opcional
    pdf_generado = False
    nombre_pdf = None
    pdf_final = None

    if formato in ("PDF", "AMBOS"):
        nombre_pdf = generar_nombre_documento(
            tipo="ANEXO_EXTENSION_CONTRATO",
            ap_paterno=trabajador.ap_paterno or "SIN_APELLIDO",
            fecha_ref=fecha_ref,
            extension="pdf",
        )
        pdf_final = dest_dir / nombre_pdf
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())
        pdf_generado = True

    # Persistencia en el MISMO registro
    anexo.formato = formato
    anexo.docx_nombre = nombre_docx
    anexo.docx_ruta = str(docx_final)

    if pdf_generado:
        anexo.pdf_nombre = nombre_pdf
        anexo.pdf_ruta = str(pdf_final)

    meta = anexo.meta or {}
    if isinstance(meta, dict):
        gen = meta.get("generacion", {})
        if not isinstance(gen, dict):
            gen = {}
        gen["ultimo_formato"] = formato
        gen["ultimo_ts"] = datetime.now().isoformat(timespec="seconds")
        gen["ultimo_resultado"] = "OK"
        meta["generacion"] = gen
        anexo.meta = meta


# ==========================
# LISTADO (solo nuevos)
# ==========================
@bp.get("/")
@login_required
@role_required("ADMIN", "OPERADOR")
def anexos_listado():
    trabajador_id = request.args.get("trabajador_id", type=int)
    contrato_id = request.args.get("contrato_id", type=int)

    q = AnexoExtensionContrato.query.order_by(
        AnexoExtensionContrato.fecha_anexo.desc(),
        AnexoExtensionContrato.id.desc(),
    )

    if trabajador_id:
        q = q.filter(AnexoExtensionContrato.trabajador_id == trabajador_id)
    if contrato_id:
        q = q.filter(AnexoExtensionContrato.contrato_id == contrato_id)

    anexos_nuevos = q.all()

    return render_template(
        "anexos_listado.html",
        anexos_nuevos=anexos_nuevos,
        filtro_trabajador_id=trabajador_id,
        filtro_contrato_id=contrato_id,
    )


# ==========================
# NUEVO: ANEXO EXTENSIÓN (GET/POST)
# ==========================
@bp.route("/contratos/<int:contrato_id>/extension", methods=["GET", "POST"])
@login_required
@role_required("ADMIN", "OPERADOR")
def anexo_extension(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    if request.method == "GET":
        return render_template(
            "anexos/anexo_extension_nuevo.html",
            contrato=contrato,
            fecha_termino_anterior=getattr(contrato, "fecha_termino", None),
        )

    # POST
    formato = (request.form.get("formato") or "AMBOS").upper().strip()
    if formato not in ("DOCX", "PDF", "AMBOS"):
        formato = "AMBOS"

    fecha_anexo = parse_date(request.form.get("fecha_anexo")) or date.today()
    nueva_fecha_termino = parse_date(request.form.get("fecha_termino_nueva"))
    observaciones = (request.form.get("observaciones") or "").strip() or None

    if not nueva_fecha_termino:
        flash("Debes indicar una fecha de término nueva válida.", "danger")
        return redirect(url_for("anexos.anexo_extension", contrato_id=contrato_id))

    anterior = getattr(contrato, "fecha_termino", None)

    if not contrato.trabajador:
        flash("Contrato sin trabajador asociado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    if not contrato.empleador:
        flash("El contrato no tiene empleador asignado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    if not getattr(contrato, "obra", None) or not (getattr(contrato.obra, "centro_costo", "") or "").strip():
        flash("El contrato no tiene obra con centro de costo definido. No se puede generar en Nextcloud.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    # DEDUPE DEFENSIVO
    existente = (
        AnexoExtensionContrato.query
        .filter_by(
            contrato_id=contrato.id,
            fecha_anexo=fecha_anexo,
            fecha_termino_nueva=nueva_fecha_termino,
        )
        .order_by(AnexoExtensionContrato.id.desc())
        .first()
    )

    if existente:
        flash(
            f"Este anexo ya existía (ID #{existente.id}). "
            f"No se creó un nuevo registro; puedes regenerar archivos si lo necesitas.",
            "info",
        )
        return redirect(url_for("anexos.anexo_extension_generar", anexo_id=existente.id))

    # Creamos registro (sin archivos aún)
    anexo = AnexoExtensionContrato(
        contrato_id=contrato.id,
        trabajador_id=contrato.trabajador.id,
        empleador_id=contrato.empleador_id,
        obra_id=contrato.obra_id,
        fecha_anexo=fecha_anexo,
        fecha_termino_anterior=anterior,
        fecha_termino_nueva=nueva_fecha_termino,
        observaciones=observaciones,
        formato=formato,
        estado="VIGENTE",
        docx_nombre=None,
        docx_ruta=None,
        pdf_nombre=None,
        pdf_ruta=None,
        meta={"tipo": "EXTENSION_CONTRATO", "generacion": {"ultimo_resultado": None}},
    )
    db.session.add(anexo)

    # Actualizar contrato al nuevo término
    contrato.fecha_termino = nueva_fecha_termino

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo crear el anexo en BD: {e}", "danger")
        return redirect(url_for("anexos.anexo_extension", contrato_id=contrato.id))

    # Generación inicial de archivos
    try:
        _generar_archivos_anexo_extension_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(
            f"El anexo se creó en BD (ID #{anexo.id}), pero falló la generación de archivos: {e}. "
            f"Puedes reintentar desde 'Regenerar archivos (Nextcloud)'.",
            "warning",
        )
        return redirect(url_for("anexos.anexo_extension_detalle", anexo_id=anexo.id))

    flash("Anexo de extensión creado (BD) y archivos generados en Nextcloud.", "success")
    return redirect(url_for("anexos.anexo_extension_detalle", anexo_id=anexo.id))


# ==========================
# GENERAR / REGENERAR ARCHIVOS (sin BD nueva)
# ==========================
@bp.get("/extension/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def anexo_extension_generar(anexo_id: int):
    anexo = AnexoExtensionContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    return render_template(
        "anexos/anexo_extension_generar.html",
        anexo=anexo,
        contrato=contrato,
        formato_default=(anexo.formato or "AMBOS").upper(),
    )


@bp.post("/extension/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def anexo_extension_generar_post(anexo_id: int):
    anexo = AnexoExtensionContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    formato = (request.form.get("formato") or anexo.formato or "AMBOS").upper().strip()
    if formato not in ("DOCX", "PDF", "AMBOS"):
        flash("Formato inválido. Usa DOCX, PDF o AMBOS.", "danger")
        return redirect(url_for("anexos.anexo_extension_generar", anexo_id=anexo.id))

    try:
        _generar_archivos_anexo_extension_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo regenerar el anexo en Nextcloud: {e}", "danger")
        return redirect(url_for("anexos.anexo_extension_generar", anexo_id=anexo.id))

    flash("Archivos regenerados en Nextcloud (sin duplicar registro en BD).", "success")
    return redirect(url_for("anexos.anexo_extension_detalle", anexo_id=anexo.id))


# ==========================
# DETALLE (tabla nueva)
# ==========================
@bp.get("/extension/<int:anexo_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def anexo_extension_detalle(anexo_id: int):
    anexo = AnexoExtensionContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    return render_template("anexos/anexo_extension_detalle.html", anexo=anexo)


# ==========================
# EDITAR (OPERADOR + ADMIN)
# ==========================
@bp.route("/extension/<int:anexo_id>/editar", methods=["GET", "POST"])
@login_required
@role_required("ADMIN", "OPERADOR")
def anexo_extension_editar(anexo_id: int):
    anexo = AnexoExtensionContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    if request.method == "GET":
        return render_template("anexos/anexo_extension_editar.html", anexo=anexo)

    formato = (request.form.get("formato") or anexo.formato or "AMBOS").upper().strip()
    if formato not in ("DOCX", "PDF", "AMBOS"):
        formato = anexo.formato or "AMBOS"

    estado = (request.form.get("estado") or anexo.estado or "VIGENTE").upper().strip()
    if estado not in ("VIGENTE", "ANULADO", "REEMPLAZADO"):
        estado = anexo.estado or "VIGENTE"

    fecha_anexo = parse_date(request.form.get("fecha_anexo")) or anexo.fecha_anexo or date.today()
    fecha_termino_nueva = parse_date(request.form.get("fecha_termino_nueva")) or anexo.fecha_termino_nueva
    observaciones = (request.form.get("observaciones") or "").strip() or None

    if not fecha_termino_nueva:
        flash("Debes indicar una nueva fecha de término válida.", "danger")
        return redirect(url_for("anexos.anexo_extension_editar", anexo_id=anexo.id))

    try:
        anexo.formato = formato
        anexo.estado = estado
        anexo.fecha_anexo = fecha_anexo
        anexo.fecha_termino_nueva = fecha_termino_nueva
        anexo.observaciones = observaciones

        db.session.commit()
        flash("Anexo actualizado correctamente.", "success")
        return redirect(url_for("anexos.anexo_extension_detalle", anexo_id=anexo.id))

    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo actualizar el anexo: {e}", "danger")
        return redirect(url_for("anexos.anexo_extension_editar", anexo_id=anexo.id))


# ==========================
# ELIMINAR (solo ADMIN)
# ==========================
@bp.post("/extension/<int:anexo_id>/eliminar")
@login_required
@role_required("ADMIN")
def anexo_extension_eliminar(anexo_id: int):
    anexo = AnexoExtensionContrato.query.get_or_404(anexo_id)

    try:
        db.session.delete(anexo)
        db.session.commit()
        flash("Anexo eliminado de la base de datos.", "warning")
        return redirect(url_for("anexos.anexos_listado"))

    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo eliminar el anexo: {e}", "danger")
        return redirect(url_for("anexos.anexo_extension_detalle", anexo_id=anexo.id))


# ==========================
# DETALLE ANEXO (EventoLaboral) - legacy
# ==========================
@bp.get("/<int:evento_id>")
@login_required
@role_required("ADMIN")
def anexo_detalle(evento_id: int):
    evento = EventoLaboral.query.get_or_404(evento_id)
    if evento.categoria != "ANEXO":
        flash("El evento solicitado no corresponde a un anexo.", "warning")
        return redirect(url_for("anexos.anexos_listado"))

    return render_template("anexo_detalle.html", anexo=evento)
