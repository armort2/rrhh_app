# web/app/blueprints/anexos_cambio_cargo/routes.py

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from docx import Document
from flask import abort, current_app, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required

from ...auth.decorators import role_required
from ...extensions import db
from ...models import AnexoCambioCargoContrato, Cargo, Contrato, EventoLaboral
from ...services.docx_engine import insert_text_block_at_placeholder, replace_placeholders
from ...services.paths_nextcloud import get_nc_paths
from ...services.pdf_convert import convert_docx_to_pdf
from ...services.storage_nextcloud_fs import ensure_dir
from ...utils import fecha_larga_es, format_rut_parts, parse_date
from . import bp


# ==========================
# Helpers generales
# ==========================

def _require_contrato_access(contrato: Contrato) -> None:
    """
    Control de acceso por obra asociada al contrato.
    - ADMIN: acceso total
    - Otros: solo si la obra del contrato está asignada al usuario.
    """
    if current_user.has_role("ADMIN"):
        return

    obra_id = getattr(contrato, "obra_id", None)
    if obra_id is None:
        abort(403)

    if not current_user.can_access_obra(int(obra_id)):
        abort(403)


def _norm_formato(value: str | None, default: str = "AMBOS") -> str:
    fmt = (value or default).strip().upper()
    return fmt if fmt in ("DOCX", "PDF", "AMBOS") else default


def _format_rut_documento(rut_value, dv_value=None) -> str:
    """
    Formatea RUT para documentos laborales.

    Casos soportados:
    - rut + dv separados: 11546804 + 9 -> 11.546.804-9
    - rut completo con guion: 11546804-9 / 11.546.804-9 -> 11.546.804-9
    - si viene solo el cuerpo sin DV, lo devuelve intacto para no inventar datos.
    """
    rut_raw = (str(rut_value).strip().upper() if rut_value is not None else "")
    dv_raw = (str(dv_value).strip().upper() if dv_value is not None else "")

    if rut_raw and dv_raw:
        return format_rut_parts(rut_raw, dv_raw) or f"{rut_raw}-{dv_raw}"

    if not rut_raw:
        return ""

    if "-" in rut_raw:
        clean = re.sub(r"[^0-9K]", "", rut_raw)
        if len(clean) >= 2:
            cuerpo = clean[:-1]
            dv = clean[-1]
            return format_rut_parts(cuerpo, dv) or rut_raw

    return rut_raw


def _safe_filename(name: str) -> str:
    name = (name or "").strip()
    name = re.sub(r'[\\\/:\*\?"<>\|]+', " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name.rstrip(" .")


def _nombre_anexo_cambio_cargo(trabajador, fecha_ref: date, extension: str) -> str:
    ap_paterno = (getattr(trabajador, "ap_paterno", "") or "").strip() or "SIN_APELLIDO"
    ap_materno = (getattr(trabajador, "ap_materno", "") or "").strip()
    nombres = (getattr(trabajador, "nombres", "") or "").strip()
    nombre_persona = " ".join([x for x in [ap_paterno, ap_materno, nombres] if x]).strip()
    base = _safe_filename(f"{fecha_ref.strftime('%Y-%m-%d')} Anexo cambio cargo - {nombre_persona}")
    ext = (extension or "").lstrip(".").lower()
    return f"{base}.{ext}"


def _require_obra_y_centro_costo(contrato: Contrato) -> str:
    obra = getattr(contrato, "obra", None)
    if not obra:
        raise RuntimeError("El contrato no tiene obra asignada (no se puede determinar centro de costo).")

    centro_costo = (getattr(obra, "centro_costo", "") or "").strip()
    if not centro_costo:
        raise RuntimeError(f"La obra '{getattr(obra, 'nombre', 'SIN_NOMBRE')}' no tiene centro de costo definido.")
    return centro_costo


# ==========================
# DOCX builder
# ==========================

def _get_tpl_anexo_cambio_cargo() -> Path:
    return Path(current_app.root_path) / "document_templates" / "anexos" / "anexo_cambio_cargo.docx"


def _get_tpl_cargos_library() -> Path:
    return Path(current_app.root_path) / "document_templates" / "contratos" / "cargos.docx"


def _build_variables_anexo_cambio_cargo(anexo: AnexoCambioCargoContrato, contrato: Contrato) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    trabajador_rut = _format_rut_documento(
        getattr(trabajador, "rut", ""),
        getattr(trabajador, "dv", ""),
    )

    empleador_rut = _format_rut_documento(
        getattr(empleador, "rut", "") if empleador else "",
    )

    rep_legal_rut = _format_rut_documento(
        getattr(empleador, "rut_rep_legal", "") if empleador else "",
    )

    return {
        "{{OBRA_NOMBRE}}": (obra.nombre if obra else ""),
        "{{OBRA_COMUNA}}": (obra.comuna if obra else ""),
        "{{ANEXO_CAMBIO_CARGO_FECHA}}": fecha_larga_es(anexo.fecha_cambio or date.today()),

        "{{EMPLEADOR_RAZON_SOCIAL}}": (empleador.razon_social if empleador else ""),
        "{{EMPLEADOR_RUT}}": empleador_rut,
        "{{REP_LEGAL_NOMBRE}}": (getattr(empleador, "nombre_rep_legal", "") or "") if empleador else "",
        "{{REP_LEGAL_RUT}}": rep_legal_rut,
        "{{EMPLEADOR_DIRECCION}}": (empleador.direccion or "") if empleador else "",
        "{{EMPLEADOR_COMUNA}}": (empleador.comuna or "") if empleador else "",

        "{{TRABAJADOR_NOMBRES}}": (trabajador.nombres or ""),
        "{{TRABAJADOR_AP_PATERNO}}": (trabajador.ap_paterno or ""),
        "{{TRABAJADOR_AP_MATERNO}}": (trabajador.ap_materno or ""),
        "{{TRABAJADOR_RUT}}": trabajador_rut,

        "{{CARGO_ANTERIOR_NOMBRE}}": anexo.cargo_anterior_nombre or "",
        "{{CARGO_NUEVO_NOMBRE}}": anexo.cargo_nuevo_nombre or "",
    }


def _build_anexo_docx(
    *,
    template_path: Path,
    cargos_library_path: Path,
    cargo_id: int,
    variables: dict[str, str],
    output_path: Path,
) -> Path:
    """
    Genera el anexo e inserta la descripción contractual del cargo desde cargos.docx.
    Usa el mismo estándar de marcadores de contratos: CARGO_{id}__DESCRIPCION.
    """
    doc = Document(str(template_path))
    cargos_doc = Document(str(cargos_library_path))

    replace_placeholders(doc, variables)

    insert_text_block_at_placeholder(
        target_doc=doc,
        placeholder="{{CARGO_DESCRIPCION}}",
        source_doc=cargos_doc,
        bookmark_name=f"CARGO_{cargo_id}__DESCRIPCION",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _touch_meta_generacion(anexo: AnexoCambioCargoContrato, *, formato: str, resultado: str, detalle: str | None = None) -> None:
    meta = anexo.meta or {}
    if not isinstance(meta, dict):
        meta = {}

    gen = meta.get("generacion", {})
    if not isinstance(gen, dict):
        gen = {}

    gen["ultimo_formato"] = formato
    gen["ultimo_ts"] = datetime.now().isoformat(timespec="seconds")
    gen["ultimo_resultado"] = resultado
    if detalle:
        gen["detalle"] = detalle

    meta["generacion"] = gen
    anexo.meta = meta


def _generar_archivos_anexo_cambio_cargo_nextcloud(
    *,
    anexo: AnexoCambioCargoContrato,
    contrato: Contrato,
    formato: str,
) -> None:
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")
    if not anexo.cargo_nuevo_id:
        raise RuntimeError("El anexo no tiene cargo nuevo asociado.")

    centro_costo = _require_obra_y_centro_costo(contrato)
    trabajador = contrato.trabajador

    tpl = _get_tpl_anexo_cambio_cargo()
    if not tpl.exists():
        raise RuntimeError("No se encontró la plantilla 'anexo_cambio_cargo.docx' en document_templates/anexos/.")

    cargos_tpl = _get_tpl_cargos_library()
    if not cargos_tpl.exists():
        raise RuntimeError("No se encontró 'cargos.docx' en document_templates/contratos/.")

    variables = _build_variables_anexo_cambio_cargo(anexo=anexo, contrato=contrato)

    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    ensure_dir(dest_dir)
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_cambio_cargo_{contrato.id}_{anexo.id}_{stamp}.docx"

    _build_anexo_docx(
        template_path=tpl,
        cargos_library_path=cargos_tpl,
        cargo_id=int(anexo.cargo_nuevo_id),
        variables=variables,
        output_path=docx_tmp,
    )

    fecha_ref = anexo.fecha_cambio or date.today()
    nombre_docx = _nombre_anexo_cambio_cargo(trabajador=trabajador, fecha_ref=fecha_ref, extension="docx")
    docx_final = dest_dir / nombre_docx

    if formato in ("DOCX", "AMBOS"):
        docx_final.write_bytes(docx_tmp.read_bytes())
        anexo.docx_nombre = nombre_docx
        anexo.docx_ruta = str(docx_final)

    if formato in ("PDF", "AMBOS"):
        nombre_pdf = _nombre_anexo_cambio_cargo(trabajador=trabajador, fecha_ref=fecha_ref, extension="pdf")
        pdf_final = dest_dir / nombre_pdf
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())
        anexo.pdf_nombre = nombre_pdf
        anexo.pdf_ruta = str(pdf_final)

    anexo.formato = formato
    _touch_meta_generacion(anexo, formato=formato, resultado="OK")


# ==========================
# Routes
# ==========================

@bp.get("/")
@login_required
@role_required("ADMIN", "OPERADOR", "REVISOR")
def listado():
    q = AnexoCambioCargoContrato.query.order_by(
        AnexoCambioCargoContrato.fecha_cambio.desc(),
        AnexoCambioCargoContrato.id.desc(),
    )

    if not current_user.has_role("ADMIN"):
        obra_ids = [int(o.id) for o in getattr(current_user, "obras", [])]
        q = q.filter(AnexoCambioCargoContrato.obra_id.in_(obra_ids or [-1]))

    anexos = q.limit(300).all()
    return render_template("anexos_cambio_cargo/listado.html", anexos=anexos)


@bp.get("/nuevo/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def nuevo(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    if contrato.estado_contrato != "VIGENTE":
        flash("No se puede generar anexo de cambio de cargo: el contrato no está vigente.", "warning")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    cargos = Cargo.query.filter_by(activo=True).order_by(Cargo.nombre).all()
    return render_template("anexos_cambio_cargo/nuevo.html", contrato=contrato, cargos=cargos)


@bp.post("/crear/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def crear(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    formato = _norm_formato(request.form.get("formato"))
    fecha_cambio = parse_date(request.form.get("fecha_cambio")) or date.today()
    observaciones = (request.form.get("observaciones") or "").strip() or None

    try:
        cargo_nuevo_id = int(request.form.get("cargo_nuevo_id") or 0)
    except (TypeError, ValueError):
        cargo_nuevo_id = 0

    cargo_nuevo = Cargo.query.get(cargo_nuevo_id) if cargo_nuevo_id else None

    if not cargo_nuevo or not getattr(cargo_nuevo, "activo", True):
        flash("Debes seleccionar un cargo nuevo válido y activo.", "danger")
        return redirect(url_for("anexos_cambio_cargo.nuevo", contrato_id=contrato.id))

    if not contrato.trabajador:
        flash("Contrato sin trabajador asociado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))
    if not contrato.empleador:
        flash("El contrato no tiene empleador asignado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))
    if not getattr(contrato, "obra", None) or not (getattr(contrato.obra, "centro_costo", "") or "").strip():
        flash("El contrato no tiene obra con centro de costo definido. No se puede generar en Nextcloud.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    cargo_anterior = contrato.cargo
    if cargo_anterior and int(cargo_anterior.id) == int(cargo_nuevo.id):
        flash("El cargo nuevo es igual al cargo actual del contrato. No se creó anexo.", "warning")
        return redirect(url_for("anexos_cambio_cargo.nuevo", contrato_id=contrato.id))

    existente = (
        AnexoCambioCargoContrato.query
        .filter_by(
            contrato_id=contrato.id,
            fecha_cambio=fecha_cambio,
            cargo_nuevo_id=cargo_nuevo.id,
        )
        .order_by(AnexoCambioCargoContrato.id.desc())
        .first()
    )
    if existente:
        flash(
            f"Este anexo ya existía (ID #{existente.id}). No se creó un nuevo registro; puedes regenerar archivos.",
            "info",
        )
        return redirect(url_for("anexos_cambio_cargo.generar", anexo_id=existente.id))

    anexo = AnexoCambioCargoContrato(
        contrato_id=contrato.id,
        trabajador_id=contrato.trabajador.id,
        empleador_id=contrato.empleador_id,
        obra_id=contrato.obra_id,
        fecha_cambio=fecha_cambio,
        cargo_anterior_id=(cargo_anterior.id if cargo_anterior else None),
        cargo_anterior_nombre=(cargo_anterior.nombre if cargo_anterior else None),
        cargo_nuevo_id=cargo_nuevo.id,
        cargo_nuevo_nombre=cargo_nuevo.nombre,
        observaciones=observaciones,
        formato=formato,
        estado="VIGENTE",
        meta={"tipo": "CAMBIO_CARGO", "generacion": {"ultimo_resultado": None}},
    )

    # Actualizamos el contrato maestro para que quede vigente con el nuevo cargo.
    contrato.cargo_id = cargo_nuevo.id
    if contrato.estado_contrato == "VIGENTE" and contrato.trabajador:
        contrato.trabajador.cargo_id = cargo_nuevo.id

    try:
        db.session.add(anexo)
        db.session.add(contrato)
        db.session.flush()

        evento = EventoLaboral(
            trabajador_id=contrato.trabajador.id,
            contrato_id=contrato.id,
            obra_id=contrato.obra_id,
            empleador_id=contrato.empleador_id,
            categoria="ANEXO",
            tipo="CAMBIO_CARGO",
            titulo=f"Anexo cambio de cargo (Contrato #{contrato.id})",
            fecha_evento=fecha_cambio,
            estado="VIGENTE",
            meta={
                "anexo_cambio_cargo_id": anexo.id,
                "cargo_anterior_id": anexo.cargo_anterior_id,
                "cargo_anterior_nombre": anexo.cargo_anterior_nombre,
                "cargo_nuevo_id": cargo_nuevo.id,
                "cargo_nuevo_nombre": cargo_nuevo.nombre,
                "formato": formato,
            },
        )
        db.session.add(evento)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo crear el anexo en BD: {e}", "danger")
        return redirect(url_for("anexos_cambio_cargo.nuevo", contrato_id=contrato.id))

    try:
        _generar_archivos_anexo_cambio_cargo_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(
            f"El anexo se creó en BD (ID #{anexo.id}) y el contrato fue actualizado, pero falló la generación de archivos: {e}. "
            "Puedes reintentar desde 'Regenerar archivos'.",
            "warning",
        )
        return redirect(url_for("anexos_cambio_cargo.detalle", anexo_id=anexo.id))

    flash("Anexo de cambio de cargo creado, contrato actualizado y archivos generados en Nextcloud.", "success")
    return redirect(url_for("anexos_cambio_cargo.detalle", anexo_id=anexo.id))


@bp.get("/<int:anexo_id>")
@login_required
@role_required("ADMIN", "OPERADOR", "REVISOR")
def detalle(anexo_id: int):
    anexo = AnexoCambioCargoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)
    return render_template("anexos_cambio_cargo/detalle.html", anexo=anexo, contrato=contrato)


@bp.get("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar(anexo_id: int):
    anexo = AnexoCambioCargoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)
    return render_template(
        "anexos_cambio_cargo/generar.html",
        anexo=anexo,
        contrato=contrato,
        formato_default=(anexo.formato or "AMBOS").upper(),
    )


@bp.post("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar_post(anexo_id: int):
    anexo = AnexoCambioCargoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    formato = _norm_formato(request.form.get("formato"), default=(anexo.formato or "AMBOS"))

    try:
        _generar_archivos_anexo_cambio_cargo_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo regenerar el anexo en Nextcloud: {e}", "danger")
        return redirect(url_for("anexos_cambio_cargo.generar", anexo_id=anexo.id))

    flash("Archivos regenerados en Nextcloud sin duplicar registro en BD.", "success")
    return redirect(url_for("anexos_cambio_cargo.detalle", anexo_id=anexo.id))


@bp.post("/<int:anexo_id>/eliminar")
@login_required
@role_required("ADMIN")
def eliminar(anexo_id: int):
    anexo = AnexoCambioCargoContrato.query.get_or_404(anexo_id)

    try:
        db.session.delete(anexo)
        db.session.commit()
        flash("Anexo eliminado de la base de datos. Los archivos físicos no fueron eliminados.", "warning")
        return redirect(url_for("anexos_cambio_cargo.listado"))
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo eliminar el anexo: {e}", "danger")
        return redirect(url_for("anexos_cambio_cargo.detalle", anexo_id=anexo.id))
