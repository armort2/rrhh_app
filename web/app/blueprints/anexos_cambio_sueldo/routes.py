# web/app/blueprints/anexos_cambio_sueldo/routes.py

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from flask import abort, current_app, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required

from ...auth.decorators import role_required
from ...extensions import db
from ...models import AnexoCambioSueldoContrato, Contrato, EventoLaboral
from ...services.docx_engine import replace_placeholders
from ...services.num_to_words_es import clp_a_palabras
from ...services.paths_nextcloud import get_nc_paths
from ...services.pdf_convert import convert_docx_to_pdf
from ...services.storage_nextcloud_fs import ensure_dir
from ...utils import fecha_larga_es, format_rut_parts, parse_date
from . import bp


# ==========================
# Helpers generales
# ==========================

def _require_contrato_access(contrato: Contrato) -> None:
    """
    Control de acceso por obra asociada al contrato.
    - ADMIN: acceso total
    - Otros: solo si la obra del contrato está asignada al usuario
    """
    if current_user.has_role("ADMIN"):
        return

    obra_id = getattr(contrato, "obra_id", None)
    if obra_id is None:
        abort(403)

    if not current_user.can_access_obra(int(obra_id)):
        abort(403)


def _norm_formato(value: str | None, default: str = "AMBOS") -> str:
    fmt = (value or default).strip().upper()
    return fmt if fmt in ("DOCX", "PDF", "AMBOS") else default


def _parse_money(value: str | None) -> Decimal | None:
    """
    Acepta formatos habituales chilenos:
    - 500000
    - 500.000
    - $500.000
    - 500000,00
    Retorna Decimal entero para CLP.
    """
    raw = (value or "").strip()
    if not raw:
        return None

    cleaned = raw.replace("$", "").replace(" ", "")

    # Si viene con coma decimal, quitamos miles con punto y usamos coma como decimal.
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    else:
        cleaned = cleaned.replace(".", "")

    try:
        amount = Decimal(cleaned).quantize(Decimal("1"))
    except (InvalidOperation, ValueError):
        return None

    if amount < 0:
        return None
    return amount


def _fmt_clp(value) -> str:
    if value is None:
        return "0"
    n = int(Decimal(str(value)).quantize(Decimal("1")))
    return f"{n:,}".replace(",", ".")


def _format_rut_documento(rut_value, dv_value=None) -> str:
    """
    Formatea RUT para documentos laborales.

    Casos soportados:
    - rut + dv separados:
        11546804 + 9 -> 11.546.804-9
    - rut completo ya con guion:
        11546804-9 -> 11.546.804-9
        11.546.804-9 -> 11.546.804-9

    Si viene solo el cuerpo del RUT sin DV, lo devuelve tal como viene,
    para no inventar un dígito verificador ni cortar mal datos existentes.
    """
    rut_raw = (str(rut_value).strip().upper() if rut_value is not None else "")
    dv_raw = (str(dv_value).strip().upper() if dv_value is not None else "")

    if rut_raw and dv_raw:
        return format_rut_parts(rut_raw, dv_raw) or f"{rut_raw}-{dv_raw}"

    if not rut_raw:
        return ""

    # Si el RUT viene completo con guion, lo normalizamos.
    if "-" in rut_raw:
        clean = re.sub(r"[^0-9K]", "", rut_raw)
        if len(clean) >= 2:
            cuerpo = clean[:-1]
            dv = clean[-1]
            return format_rut_parts(cuerpo, dv) or rut_raw

    # Si ya venía formateado o si viene solo el cuerpo sin DV, lo dejamos intacto.
    return rut_raw


def _clp_palabras_sin_pesos(value) -> str:
    """
    La plantilla trae literalmente la palabra 'pesos' después del placeholder.
    Por eso usamos el helper existente y retiramos el sufijo para evitar 'pesos pesos'.
    """
    txt = (clp_a_palabras(value) or "").strip()
    txt = re.sub(r"\s+pesos?$", "", txt, flags=re.IGNORECASE).strip()
    return txt


def _safe_filename(name: str) -> str:
    name = (name or "").strip()
    name = re.sub(r'[\\\/:\*\?"<>\|]+', " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name.rstrip(" .")


def _nombre_anexo_cambio_sueldo(trabajador, fecha_ref: date, extension: str) -> str:
    ap_paterno = (getattr(trabajador, "ap_paterno", "") or "").strip() or "SIN_APELLIDO"
    ap_materno = (getattr(trabajador, "ap_materno", "") or "").strip()
    nombres = (getattr(trabajador, "nombres", "") or "").strip()
    nombre_persona = " ".join([x for x in [ap_paterno, ap_materno, nombres] if x]).strip()
    base = _safe_filename(f"{fecha_ref.strftime('%Y-%m-%d')} Anexo cambio sueldo - {nombre_persona}")
    ext = (extension or "").lstrip(".").lower()
    return f"{base}.{ext}"


def _require_obra_y_centro_costo(contrato: Contrato) -> str:
    obra = getattr(contrato, "obra", None)
    if not obra:
        raise RuntimeError("El contrato no tiene obra asignada (no se puede determinar centro de costo).")

    centro_costo = (getattr(obra, "centro_costo", "") or "").strip()
    if not centro_costo:
        raise RuntimeError(f"La obra '{getattr(obra, 'nombre', 'SIN_NOMBRE')}' no tiene centro de costo definido.")
    return centro_costo


# ==========================
# DOCX builder
# ==========================

def _get_tpl_anexo_cambio_sueldo() -> Path:
    return Path(current_app.root_path) / "document_templates" / "anexos" / "anexo_cambio_sueldo.docx"


def _build_anexo_docx(template_path: Path, variables: dict[str, str], output_path: Path) -> Path:
    doc = Document(str(template_path))
    replace_placeholders(doc, variables)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _build_variables_anexo_cambio_sueldo(anexo: AnexoCambioSueldoContrato, contrato: Contrato) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    trabajador_rut = _format_rut_documento(
        getattr(trabajador, "rut", ""),
        getattr(trabajador, "dv", ""),
    )

    empleador_rut = _format_rut_documento(
        getattr(empleador, "rut", "") if empleador else "",
    )

    rep_legal_rut = _format_rut_documento(
        getattr(empleador, "rut_rep_legal", "") if empleador else "",
    )

    return {
        "{{OBRA_NOMBRE}}": (obra.nombre if obra else ""),
        "{{OBRA_COMUNA}}": (obra.comuna if obra else ""),
        "{{ANEXO_CAMBIO_SUELDO_FECHA}}": fecha_larga_es(anexo.fecha_cambio or date.today()),

        "{{EMPLEADOR_RAZON_SOCIAL}}": (empleador.razon_social if empleador else ""),
        "{{EMPLEADOR_RUT}}": empleador_rut,
        "{{REP_LEGAL_NOMBRE}}": (getattr(empleador, "nombre_rep_legal", "") or "") if empleador else "",
        "{{REP_LEGAL_RUT}}": rep_legal_rut,
        "{{EMPLEADOR_DIRECCION}}": (empleador.direccion or "") if empleador else "",
        "{{EMPLEADOR_COMUNA}}": (empleador.comuna or "") if empleador else "",

        "{{TRABAJADOR_NOMBRES}}": (trabajador.nombres or ""),
        "{{TRABAJADOR_AP_PATERNO}}": (trabajador.ap_paterno or ""),
        "{{TRABAJADOR_AP_MATERNO}}": (trabajador.ap_materno or ""),
        "{{TRABAJADOR_RUT}}": trabajador_rut,

        "{{CONTRATO_SUELDO_BASE}}": _fmt_clp(anexo.sueldo_base_nuevo),
        "{{CONTRATO_SUELDO_BASE_PALABRAS}}": _clp_palabras_sin_pesos(anexo.sueldo_base_nuevo),
        "{{ASIG_COLACION}}": _fmt_clp(anexo.asignacion_colacion_nueva),
        "{{ASIG_COLACION_PALABRAS}}": _clp_palabras_sin_pesos(anexo.asignacion_colacion_nueva),
        "{{ASIG_MOVILIZACION}}": _fmt_clp(anexo.asignacion_movilizacion_nueva),
        "{{ASIG_MOVILIZACION_PALABRAS}}": _clp_palabras_sin_pesos(anexo.asignacion_movilizacion_nueva),
    }


def _touch_meta_generacion(anexo: AnexoCambioSueldoContrato, *, formato: str, resultado: str, detalle: str | None = None) -> None:
    meta = anexo.meta or {}
    if not isinstance(meta, dict):
        meta = {}

    gen = meta.get("generacion", {})
    if not isinstance(gen, dict):
        gen = {}

    gen["ultimo_formato"] = formato
    gen["ultimo_ts"] = datetime.now().isoformat(timespec="seconds")
    gen["ultimo_resultado"] = resultado
    if detalle:
        gen["detalle"] = detalle

    meta["generacion"] = gen
    anexo.meta = meta


def _generar_archivos_anexo_cambio_sueldo_nextcloud(
    *,
    anexo: AnexoCambioSueldoContrato,
    contrato: Contrato,
    formato: str,
) -> None:
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")

    centro_costo = _require_obra_y_centro_costo(contrato)
    trabajador = contrato.trabajador

    tpl = _get_tpl_anexo_cambio_sueldo()
    if not tpl.exists():
        raise RuntimeError("No se encontró la plantilla 'anexo_cambio_sueldo.docx' en document_templates/anexos/.")

    variables = _build_variables_anexo_cambio_sueldo(anexo=anexo, contrato=contrato)

    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    ensure_dir(dest_dir)
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_cambio_sueldo_{contrato.id}_{anexo.id}_{stamp}.docx"

    _build_anexo_docx(tpl, variables, docx_tmp)

    fecha_ref = anexo.fecha_cambio or date.today()
    nombre_docx = _nombre_anexo_cambio_sueldo(trabajador=trabajador, fecha_ref=fecha_ref, extension="docx")
    docx_final = dest_dir / nombre_docx
    docx_final.write_bytes(docx_tmp.read_bytes())

    pdf_generado = False
    nombre_pdf = None
    pdf_final = None

    if formato in ("PDF", "AMBOS"):
        nombre_pdf = _nombre_anexo_cambio_sueldo(trabajador=trabajador, fecha_ref=fecha_ref, extension="pdf")
        pdf_final = dest_dir / nombre_pdf
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())
        pdf_generado = True

    anexo.formato = formato
    anexo.docx_nombre = nombre_docx
    anexo.docx_ruta = str(docx_final)

    if pdf_generado:
        anexo.pdf_nombre = nombre_pdf
        anexo.pdf_ruta = str(pdf_final)

    _touch_meta_generacion(anexo, formato=formato, resultado="OK")


# ==========================
# Routes
# ==========================

@bp.get("/")
@login_required
@role_required("ADMIN", "OPERADOR", "REVISOR")
def listado():
    q = AnexoCambioSueldoContrato.query.order_by(
        AnexoCambioSueldoContrato.fecha_cambio.desc(),
        AnexoCambioSueldoContrato.id.desc(),
    )

    if not current_user.has_role("ADMIN"):
        obra_ids = [int(o.id) for o in getattr(current_user, "obras", [])]
        q = q.filter(AnexoCambioSueldoContrato.obra_id.in_(obra_ids or [-1]))

    anexos = q.limit(300).all()
    return render_template("anexos_cambio_sueldo/listado.html", anexos=anexos)


@bp.get("/nuevo/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def nuevo(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    if contrato.estado_contrato != "VIGENTE":
        flash("No se puede generar anexo de cambio de sueldo: el contrato no está vigente.", "warning")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    return render_template("anexos_cambio_sueldo/nuevo.html", contrato=contrato)


@bp.post("/crear/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def crear(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    formato = _norm_formato(request.form.get("formato"))
    fecha_cambio = parse_date(request.form.get("fecha_cambio")) or date.today()
    observaciones = (request.form.get("observaciones") or "").strip() or None

    sueldo_base_nuevo = _parse_money(request.form.get("sueldo_base_nuevo"))
    colacion_nueva = _parse_money(request.form.get("asignacion_colacion_nueva"))
    movilizacion_nueva = _parse_money(request.form.get("asignacion_movilizacion_nueva"))

    if sueldo_base_nuevo is None or colacion_nueva is None or movilizacion_nueva is None:
        flash("Debes indicar montos válidos para sueldo base, colación y movilización.", "danger")
        return redirect(url_for("anexos_cambio_sueldo.nuevo", contrato_id=contrato.id))

    if not contrato.trabajador:
        flash("Contrato sin trabajador asociado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))
    if not contrato.empleador:
        flash("El contrato no tiene empleador asignado.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))
    if not getattr(contrato, "obra", None) or not (getattr(contrato.obra, "centro_costo", "") or "").strip():
        flash("El contrato no tiene obra con centro de costo definido. No se puede generar en Nextcloud.", "danger")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    existente = (
        AnexoCambioSueldoContrato.query
        .filter_by(
            contrato_id=contrato.id,
            fecha_cambio=fecha_cambio,
            sueldo_base_nuevo=sueldo_base_nuevo,
            asignacion_colacion_nueva=colacion_nueva,
            asignacion_movilizacion_nueva=movilizacion_nueva,
        )
        .order_by(AnexoCambioSueldoContrato.id.desc())
        .first()
    )
    if existente:
        flash(
            f"Este anexo ya existía (ID #{existente.id}). No se creó un nuevo registro; puedes regenerar archivos.",
            "info",
        )
        return redirect(url_for("anexos_cambio_sueldo.generar", anexo_id=existente.id))

    anexo = AnexoCambioSueldoContrato(
        contrato_id=contrato.id,
        trabajador_id=contrato.trabajador.id,
        empleador_id=contrato.empleador_id,
        obra_id=contrato.obra_id,
        fecha_cambio=fecha_cambio,
        sueldo_base_anterior=contrato.sueldo_base,
        sueldo_base_nuevo=sueldo_base_nuevo,
        asignacion_colacion_anterior=contrato.asignacion_colacion,
        asignacion_colacion_nueva=colacion_nueva,
        asignacion_movilizacion_anterior=contrato.asignacion_movilizacion,
        asignacion_movilizacion_nueva=movilizacion_nueva,
        observaciones=observaciones,
        formato=formato,
        estado="VIGENTE",
        meta={"tipo": "CAMBIO_SUELDO", "generacion": {"ultimo_resultado": None}},
    )

    # Actualizamos el contrato para que el maestro quede vigente con las nuevas condiciones.
    contrato.sueldo_base = sueldo_base_nuevo
    contrato.asignacion_colacion = colacion_nueva
    contrato.asignacion_movilizacion = movilizacion_nueva

    try:
        db.session.add(anexo)
        db.session.add(contrato)
        db.session.flush()

        evento = EventoLaboral(
            trabajador_id=contrato.trabajador.id,
            contrato_id=contrato.id,
            obra_id=contrato.obra_id,
            empleador_id=contrato.empleador_id,
            categoria="ANEXO",
            tipo="CAMBIO_SUELDO",
            titulo=f"Anexo cambio de sueldo (Contrato #{contrato.id})",
            fecha_evento=fecha_cambio,
            estado="VIGENTE",
            meta={
                "anexo_cambio_sueldo_id": anexo.id,
                "sueldo_base_anterior": str(anexo.sueldo_base_anterior) if anexo.sueldo_base_anterior is not None else None,
                "sueldo_base_nuevo": str(sueldo_base_nuevo),
                "asignacion_colacion_anterior": str(anexo.asignacion_colacion_anterior) if anexo.asignacion_colacion_anterior is not None else None,
                "asignacion_colacion_nueva": str(colacion_nueva),
                "asignacion_movilizacion_anterior": str(anexo.asignacion_movilizacion_anterior) if anexo.asignacion_movilizacion_anterior is not None else None,
                "asignacion_movilizacion_nueva": str(movilizacion_nueva),
                "formato": formato,
            },
        )
        db.session.add(evento)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo crear el anexo en BD: {e}", "danger")
        return redirect(url_for("anexos_cambio_sueldo.nuevo", contrato_id=contrato.id))

    try:
        _generar_archivos_anexo_cambio_sueldo_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(
            f"El anexo se creó en BD (ID #{anexo.id}) y el contrato fue actualizado, pero falló la generación de archivos: {e}. "
            "Puedes reintentar desde 'Regenerar archivos'.",
            "warning",
        )
        return redirect(url_for("anexos_cambio_sueldo.detalle", anexo_id=anexo.id))

    flash("Anexo de cambio de sueldo creado, contrato actualizado y archivos generados en Nextcloud.", "success")
    return redirect(url_for("anexos_cambio_sueldo.detalle", anexo_id=anexo.id))


@bp.get("/<int:anexo_id>")
@login_required
@role_required("ADMIN", "OPERADOR", "REVISOR")
def detalle(anexo_id: int):
    anexo = AnexoCambioSueldoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)
    return render_template("anexos_cambio_sueldo/detalle.html", anexo=anexo, contrato=contrato)


@bp.get("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar(anexo_id: int):
    anexo = AnexoCambioSueldoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)
    return render_template(
        "anexos_cambio_sueldo/generar.html",
        anexo=anexo,
        contrato=contrato,
        formato_default=(anexo.formato or "AMBOS").upper(),
    )


@bp.post("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar_post(anexo_id: int):
    anexo = AnexoCambioSueldoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    formato = _norm_formato(request.form.get("formato"), default=(anexo.formato or "AMBOS"))

    try:
        _generar_archivos_anexo_cambio_sueldo_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo regenerar el anexo en Nextcloud: {e}", "danger")
        return redirect(url_for("anexos_cambio_sueldo.generar", anexo_id=anexo.id))

    flash("Archivos regenerados en Nextcloud sin duplicar registro en BD.", "success")
    return redirect(url_for("anexos_cambio_sueldo.detalle", anexo_id=anexo.id))


@bp.post("/<int:anexo_id>/eliminar")
@login_required
@role_required("ADMIN")
def eliminar(anexo_id: int):
    anexo = AnexoCambioSueldoContrato.query.get_or_404(anexo_id)

    try:
        db.session.delete(anexo)
        db.session.commit()
        flash("Anexo eliminado de la base de datos. Los archivos físicos no fueron eliminados.", "warning")
        return redirect(url_for("anexos_cambio_sueldo.listado"))
    except Exception as e:
        db.session.rollback()
        flash(f"No se pudo eliminar el anexo: {e}", "danger")
        return redirect(url_for("anexos_cambio_sueldo.detalle", anexo_id=anexo.id))