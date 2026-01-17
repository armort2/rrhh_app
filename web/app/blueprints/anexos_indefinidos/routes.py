# web/app/blueprints/anexos_indefinidos/routes.py

from __future__ import annotations

from datetime import date, timedelta, datetime
from pathlib import Path

from flask import current_app, render_template, request, redirect, url_for, flash, abort
from flask_login import login_required, current_user
from sqlalchemy.exc import IntegrityError

from docx import Document

from ...extensions import db
from ...auth.decorators import role_required
from ...config import generar_nombre_documento
from ...models import (
    Contrato,
    AnexoExtensionContrato,
    AnexoIndefinidoContrato,
)
from ...services.docx_engine import replace_placeholders
from ...services.paths_nextcloud import get_nc_paths
from ...services.pdf_convert import convert_docx_to_pdf
from ...services.storage_nextcloud_fs import ensure_dir
from ...utils import fecha_larga_es

from . import bp


# ==========================
# Helpers
# ==========================

def _parse_date(value: str | None) -> date | None:
    """
    Parse robusto para <input type="date"> (YYYY-MM-DD).
    """
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except Exception:
        return None


def _norm_tipo_contrato(value: str | None) -> str:
    raw = (value or "").strip().upper()
    raw = raw.replace("-", " ").replace("_", " ")
    raw = " ".join(raw.split())
    return raw


def _require_contrato_access(contrato: Contrato) -> None:
    """
    Control de acceso por obra asociada al contrato.
    - ADMIN: acceso total
    - Otros: solo si la obra del contrato está asignada al usuario
    """
    if current_user.has_role("ADMIN"):
        return

    obra_id = contrato.obra_id
    if obra_id is None:
        abort(403)

    if not current_user.can_access_obra(int(obra_id)):
        abort(403)


def _get_fecha_termino_referencial_y_sugerencia_desde(contrato: Contrato):
    """
    - Si hay anexo extensión vigente => usar fecha_termino_nueva como referencial
    - Si no => usar contrato.fecha_termino
    Sugerencia fecha_indefinido_desde = referencial + 1 día (si existe), si no => hoy
    """
    anexo_ext = (
        AnexoExtensionContrato.query
        .filter_by(contrato_id=contrato.id, estado="VIGENTE")
        .order_by(AnexoExtensionContrato.fecha_anexo.desc(), AnexoExtensionContrato.id.desc())
        .first()
    )

    fecha_ref = None
    if anexo_ext and anexo_ext.fecha_termino_nueva:
        fecha_ref = anexo_ext.fecha_termino_nueva
    elif contrato.fecha_termino:
        fecha_ref = contrato.fecha_termino

    if fecha_ref:
        return fecha_ref, (fecha_ref + timedelta(days=1))
    return None, date.today()


def _tiene_extension(contrato_id: int) -> bool:
    return (
        db.session.query(AnexoExtensionContrato.id)
        .filter_by(contrato_id=contrato_id)
        .first()
        is not None
    )


def _es_plazo_fijo(contrato: Contrato) -> bool:
    tipo = _norm_tipo_contrato(contrato.tipo_contrato)
    return tipo in {"PLAZO FIJO", "PLAZO_FIJO"}


def _norm_formato(value: str | None, default: str = "AMBOS") -> str:
    fmt = (value or default).strip().upper()
    if fmt not in ("DOCX", "PDF", "AMBOS"):
        return default
    return fmt


def _require_obra_y_centro_costo(contrato: Contrato) -> str:
    """
    Garantiza que exista contrato.obra y que la obra tenga centro_costo.
    Retorna centro_costo limpio.
    """
    obra = getattr(contrato, "obra", None)
    if not obra:
        raise RuntimeError("El contrato no tiene obra asignada (no se puede determinar centro de costo).")

    centro_costo = (getattr(obra, "centro_costo", "") or "").strip()
    if not centro_costo:
        raise RuntimeError(f"La obra '{getattr(obra, 'nombre', 'SIN_NOMBRE')}' no tiene centro de costo definido.")

    return centro_costo


# ==========================
# DOCX builder: Anexo indefinido
# ==========================
def _get_tpl_anexo_indefinido() -> Path:
    templates_root = Path(current_app.root_path) / "document_templates" / "anexos"
    return templates_root / "anexo_indefinido_contrato.docx"


def _build_variables_anexo_indefinido(anexo: AnexoIndefinidoContrato, contrato: Contrato) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    rut_digits = (getattr(trabajador, "rut", "") or "").strip()
    dv = (getattr(trabajador, "dv", "") or "").strip().upper()
    rut_completo = f"{rut_digits}-{dv}" if rut_digits and dv else rut_digits

    # Para la cláusula "tenía como fecha de término"
    termino_anterior = anexo.fecha_termino_referencial or contrato.fecha_termino

    variables: dict[str, str] = {
        "{{OBRA_COMUNA}}": (obra.comuna if obra and obra.comuna else ""),
        "{{ANEXO_FECHA}}": fecha_larga_es(anexo.fecha_anexo or date.today()),

        "{{EMPLEADOR_RAZON_SOCIAL}}": (empleador.razon_social if empleador else ""),
        "{{EMPLEADOR_RUT}}": (getattr(empleador, "rut", "") or "") if empleador else "",
        "{{REP_LEGAL_NOMBRE}}": (getattr(empleador, "nombre_rep_legal", "") or "") if empleador else "",
        "{{REP_LEGAL_RUT}}": (getattr(empleador, "rut_rep_legal", "") or "") if empleador else "",
        "{{EMPLEADOR_DIRECCION}}": (empleador.direccion or "") if empleador else "",
        "{{EMPLEADOR_COMUNA}}": (empleador.comuna or "") if empleador else "",
        "{{OBRA_NOMBRE}}": (obra.nombre if obra else ""),

        "{{TRABAJADOR_NOMBRES}}": (trabajador.nombres or ""),
        "{{TRABAJADOR_AP_PATERNO}}": (trabajador.ap_paterno or ""),
        "{{TRABAJADOR_AP_MATERNO}}": (trabajador.ap_materno or ""),
        "{{TRABAJADOR_RUT}}": rut_completo,

        "{{CONTRATO_FECHA_INICIO}}": fecha_larga_es(contrato.fecha_inicio) if contrato.fecha_inicio else "",
        "{{CONTRATO_FECHA_TERMINO_ANTERIOR}}": fecha_larga_es(termino_anterior) if termino_anterior else "",
        "{{FECHA_INDEFINIDO_DESDE}}": fecha_larga_es(anexo.fecha_indefinido_desde) if anexo.fecha_indefinido_desde else "",
        "{{CARGO_NOMBRE}}": (contrato.cargo.nombre if contrato.cargo else ""),
        "{{ANEXO_OBSERVACIONES}}": (anexo.observaciones or ""),
    }
    return variables


def _build_anexo_indefinido_docx(template_path: Path, variables: dict[str, str], output_path: Path) -> Path:
    doc = Document(str(template_path))
    replace_placeholders(doc, variables)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _touch_meta_generacion(anexo: AnexoIndefinidoContrato, *, formato: str, resultado: str, detalle: str | None = None) -> None:
    """
    Actualiza meta.generacion (sin asumir estructura).
    resultado: "OK" o "ERROR"
    """
    meta = anexo.meta or {}
    if not isinstance(meta, dict):
        meta = {}

    gen = meta.get("generacion", {})
    if not isinstance(gen, dict):
        gen = {}

    gen["ultimo_formato"] = formato
    gen["ultimo_ts"] = datetime.now().isoformat(timespec="seconds")
    gen["ultimo_resultado"] = resultado
    if detalle:
        gen["detalle"] = detalle

    meta["generacion"] = gen
    anexo.meta = meta


def _generar_archivos_anexo_indefinido_nextcloud(
    anexo: AnexoIndefinidoContrato,
    contrato: Contrato,
    formato: str,
) -> None:
    """
    Genera/Regenera archivos (DOCX/PDF) en Nextcloud para un anexo indefinido existente.
    Importante: NO crea registros; solo actualiza docx/pdf_* del mismo anexo.
    """
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")

    # ✅ Centro de costo desde Obra (ruta legacy)
    centro_costo = _require_obra_y_centro_costo(contrato)

    trabajador = contrato.trabajador

    tpl = _get_tpl_anexo_indefinido()
    if not tpl.exists():
        raise RuntimeError(
            "No se encontró la plantilla 'anexo_indefinido_contrato.docx' en 'document_templates/anexos/'."
        )

    variables = _build_variables_anexo_indefinido(anexo=anexo, contrato=contrato)

    # ✅ Destino Nextcloud con paths_nextcloud.py nuevo
    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_indefinido_{contrato.id}_{anexo.id}_{stamp}.docx"
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)

    ensure_dir(dest_dir)

    # Construye DOCX
    _build_anexo_indefinido_docx(tpl, variables, docx_tmp)

    # Nombres finales (regeneración sobreescribe)
    fecha_ref = anexo.fecha_anexo or date.today()

    nombre_docx = generar_nombre_documento(
        tipo="ANEXO_INDEFINIDO_CONTRATO",
        ap_paterno=trabajador.ap_paterno or "SIN_APELLIDO",
        fecha_ref=fecha_ref,
        extension="docx",
    )
    docx_final = dest_dir / nombre_docx
    docx_final.write_bytes(docx_tmp.read_bytes())

    # PDF opcional
    pdf_generado = False
    nombre_pdf = None
    pdf_final = None

    if formato in ("PDF", "AMBOS"):
        nombre_pdf = generar_nombre_documento(
            tipo="ANEXO_INDEFINIDO_CONTRATO",
            ap_paterno=trabajador.ap_paterno or "SIN_APELLIDO",
            fecha_ref=fecha_ref,
            extension="pdf",
        )
        pdf_final = dest_dir / nombre_pdf
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())
        pdf_generado = True

    # Persistencia en el MISMO registro
    anexo.formato = formato
    anexo.docx_nombre = nombre_docx
    anexo.docx_ruta = str(docx_final)

    if pdf_generado:
        anexo.pdf_nombre = nombre_pdf
        anexo.pdf_ruta = str(pdf_final)
    # Si pidieron solo DOCX, no borramos PDF previo (queda como último conocido)


# ==========================
# Routes
# ==========================

@bp.get("/contrato/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def listado_por_contrato(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    anexos = (
        AnexoIndefinidoContrato.query
        .filter_by(contrato_id=contrato.id)
        .order_by(AnexoIndefinidoContrato.fecha_anexo.desc(), AnexoIndefinidoContrato.id.desc())
        .all()
    )

    return render_template(
        "anexos_indefinidos/listado.html",
        contrato=contrato,
        anexos=anexos,
    )


@bp.get("/nuevo/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def nuevo(contrato_id: int):
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    if contrato.estado_contrato != "VIGENTE":
        flash("No se puede generar anexo indefinido: el contrato no está vigente.", "warning")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    ext = _tiene_extension(contrato.id)
    if not _es_plazo_fijo(contrato) and not ext:
        flash(
            "No se puede generar anexo indefinido: el contrato no es a plazo fijo y no registra anexos de extensión.",
            "warning",
        )
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    fecha_ref, fecha_sugerida = _get_fecha_termino_referencial_y_sugerencia_desde(contrato)

    return render_template(
        "anexos_indefinidos/nuevo.html",
        contrato=contrato,
        fecha_termino_referencial=fecha_ref,
        fecha_indefinido_sugerida=fecha_sugerida,
    )


@bp.post("/crear/<int:contrato_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def crear(contrato_id: int):
    """
    Crea el registro en BD UNA vez y redirige a pantalla de generación (Nextcloud).
    Evita duplicados por doble envío mediante PRG + dedupe defensivo.
    """
    contrato = Contrato.query.get_or_404(contrato_id)
    _require_contrato_access(contrato)

    if contrato.estado_contrato != "VIGENTE":
        flash("No se puede crear anexo indefinido: el contrato no está vigente.", "warning")
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    ext = _tiene_extension(contrato.id)
    if not _es_plazo_fijo(contrato) and not ext:
        flash(
            "No se puede crear anexo indefinido: el contrato no es a plazo fijo y no registra anexos de extensión.",
            "danger",
        )
        return redirect(url_for("contratos.contrato_detalle", contrato_id=contrato.id))

    fecha_anexo = _parse_date(request.form.get("fecha_anexo")) or date.today()
    fecha_indefinido_desde = _parse_date(request.form.get("fecha_indefinido_desde"))
    fecha_termino_referencial = _parse_date(request.form.get("fecha_termino_referencial"))
    observaciones = (request.form.get("observaciones") or "").strip()
    formato = _norm_formato(request.form.get("formato"), default="AMBOS")

    if not fecha_indefinido_desde:
        flash("Debes indicar la fecha desde la cual el contrato pasa a indefinido.", "danger")
        return redirect(url_for("anexos_indefinidos.nuevo", contrato_id=contrato.id))

    if fecha_termino_referencial is None:
        fecha_ref_calc, _ = _get_fecha_termino_referencial_y_sugerencia_desde(contrato)
        fecha_termino_referencial = fecha_ref_calc

    existente = (
        AnexoIndefinidoContrato.query
        .filter_by(
            contrato_id=contrato.id,
            fecha_anexo=fecha_anexo,
            fecha_indefinido_desde=fecha_indefinido_desde,
        )
        .order_by(AnexoIndefinidoContrato.id.desc())
        .first()
    )
    if existente:
        flash(
            f"Ya existía un anexo indefinido con la misma fecha de anexo y 'desde'. "
            f"Usaremos el registro existente (ID #{existente.id}).",
            "info",
        )
        return redirect(url_for("anexos_indefinidos.generar", anexo_id=existente.id))

    # Evitar múltiples anexos indefinidos vigentes (marca anteriores como REEMPLAZADO)
    anexos_vigentes = (
        AnexoIndefinidoContrato.query
        .filter_by(contrato_id=contrato.id, estado="VIGENTE")
        .order_by(AnexoIndefinidoContrato.fecha_anexo.desc(), AnexoIndefinidoContrato.id.desc())
        .all()
    )
    for a in anexos_vigentes:
        a.estado = "REEMPLAZADO"

    anexo = AnexoIndefinidoContrato(
        contrato_id=contrato.id,
        trabajador_id=contrato.trabajador_id,
        empleador_id=contrato.empleador_id,
        obra_id=contrato.obra_id,
        fecha_anexo=fecha_anexo,
        fecha_indefinido_desde=fecha_indefinido_desde,
        fecha_termino_referencial=fecha_termino_referencial,
        observaciones=observaciones or None,
        formato=formato,
        estado="VIGENTE",
        docx_nombre=None,
        docx_ruta=None,
        pdf_nombre=None,
        pdf_ruta=None,
        meta={
            "origen": "ANEXO_INDEFINIDO",
            "contrato_tipo_anterior": contrato.tipo_contrato,
            "tiene_extension": ext,
            "generacion": {
                "ultimo_resultado": None,
                "ultimo_formato": None,
                "ultimo_ts": None,
            },
        },
    )

    # Efecto en contrato (según tu negocio)
    contrato.tipo_contrato = "INDEFINIDO"
    contrato.estado_contrato = "VIGENTE"

    # ✅ FIX CRÍTICO: al pasar a INDEFINIDO, el contrato NO debe mantener fecha_termino pasada
    contrato.fecha_termino = None

    # ✅ (Opcional recomendado) limpiar campos de término si existían
    if hasattr(contrato, "causal_termino"):
        contrato.causal_termino = None
    if hasattr(contrato, "fecha_finiquito"):
        contrato.fecha_finiquito = None

    db.session.add(anexo)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash("No se pudo crear el anexo indefinido. Revisa los datos e intenta nuevamente.", "danger")
        return redirect(url_for("anexos_indefinidos.nuevo", contrato_id=contrato.id))

    flash("Registro del anexo creado. Ahora genera los archivos para Nextcloud.", "success")
    return redirect(url_for("anexos_indefinidos.generar", anexo_id=anexo.id))


@bp.get("/<int:anexo_id>")
@login_required
@role_required("ADMIN", "OPERADOR")
def detalle(anexo_id: int):
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    return render_template(
        "anexos_indefinidos/detalle.html",
        anexo=anexo,
        contrato=contrato,
    )


@bp.get("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar(anexo_id: int):
    """
    Pantalla para generar/regenerar archivos (Nextcloud) SIN crear registros nuevos.
    """
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    return render_template(
        "anexos_indefinidos/generar.html",
        anexo=anexo,
        contrato=contrato,
        formato_default=_norm_formato(anexo.formato, default="AMBOS"),
    )


@bp.post("/<int:anexo_id>/generar")
@login_required
@role_required("ADMIN", "OPERADOR")
def generar_post(anexo_id: int):
    """
    Genera/regenera archivos en Nextcloud y actualiza rutas/nombres en el MISMO registro.
    """
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    formato = _norm_formato(request.form.get("formato"), default=(anexo.formato or "AMBOS"))

    _touch_meta_generacion(anexo, formato=formato, resultado="ERROR", detalle="Inicio generación (pendiente)")

    try:
        _generar_archivos_anexo_indefinido_nextcloud(anexo=anexo, contrato=contrato, formato=formato)
        _touch_meta_generacion(anexo, formato=formato, resultado="OK", detalle=None)
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        try:
            _touch_meta_generacion(anexo, formato=formato, resultado="ERROR", detalle=str(e))
            db.session.commit()
        except Exception:
            db.session.rollback()

        flash(f"Falló la generación/regeneración de archivos: {e}", "danger")
        return redirect(url_for("anexos_indefinidos.generar", anexo_id=anexo.id))

    flash("Archivos generados/regenerados en Nextcloud (sin duplicar registros).", "success")
    return redirect(url_for("anexos_indefinidos.detalle", anexo_id=anexo.id))


@bp.get("/<int:anexo_id>/editar")
@login_required
@role_required("ADMIN", "OPERADOR")
def editar(anexo_id: int):
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    return render_template(
        "anexos_indefinidos/editar.html",
        anexo=anexo,
        contrato=contrato,
    )


@bp.post("/<int:anexo_id>/editar")
@login_required
@role_required("ADMIN", "OPERADOR")
def editar_post(anexo_id: int):
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    fecha_anexo_nueva = _parse_date(request.form.get("fecha_anexo"))
    fecha_indef_desde_nueva = _parse_date(request.form.get("fecha_indefinido_desde"))

    if fecha_anexo_nueva:
        anexo.fecha_anexo = fecha_anexo_nueva

    if fecha_indef_desde_nueva:
        anexo.fecha_indefinido_desde = fecha_indef_desde_nueva

    if "fecha_termino_referencial" in request.form:
        anexo.fecha_termino_referencial = _parse_date(request.form.get("fecha_termino_referencial"))

    anexo.observaciones = (request.form.get("observaciones") or "").strip() or None
    anexo.formato = _norm_formato(request.form.get("formato"), default=(anexo.formato or "AMBOS"))

    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash("No se pudo actualizar el anexo indefinido. Intenta nuevamente.", "danger")
        return redirect(url_for("anexos_indefinidos.editar", anexo_id=anexo.id))

    flash("Anexo indefinido actualizado.", "success")
    return redirect(url_for("anexos_indefinidos.detalle", anexo_id=anexo.id))


@bp.post("/<int:anexo_id>/eliminar")
@login_required
@role_required("ADMIN")
def eliminar(anexo_id: int):
    anexo = AnexoIndefinidoContrato.query.get_or_404(anexo_id)
    contrato = Contrato.query.get_or_404(anexo.contrato_id)
    _require_contrato_access(contrato)

    try:
        db.session.delete(anexo)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash("No se pudo eliminar el anexo indefinido. Intenta nuevamente.", "danger")
        return redirect(url_for("anexos_indefinidos.detalle", anexo_id=anexo.id))

    flash("Anexo indefinido eliminado de la base de datos. (No borra archivos)", "success")
    return redirect(url_for("anexos_indefinidos.listado_por_contrato", contrato_id=contrato.id))

@bp.get("/")
@login_required
@role_required("ADMIN", "OPERADOR")
def listado_global():
    """
    Listado global (navegable desde menú) de anexos indefinidos.
    - ADMIN: ve todo
    - OPERADOR: ve solo contratos cuyas obras pueda acceder
    """
    # Traemos un set acotado y ordenado (ERP-friendly)
    anexos_q = (
        AnexoIndefinidoContrato.query
        .order_by(AnexoIndefinidoContrato.fecha_anexo.desc(), AnexoIndefinidoContrato.id.desc())
        .limit(300)
        .all()
    )

    if current_user.has_role("ADMIN"):
        anexos = anexos_q
    else:
        # Filtrado defensivo sin asumir helpers adicionales del user model
        anexos = []
        for a in anexos_q:
            # Si no tiene obra asociada, no mostrar a no-admin
            if not a.obra_id:
                continue
            try:
                if current_user.can_access_obra(int(a.obra_id)):
                    anexos.append(a)
            except Exception:
                # Si por cualquier motivo no puede validar acceso, mejor no mostrar
                continue

    return render_template(
        "anexos_indefinidos/listado_global.html",
        anexos=anexos,
    )
