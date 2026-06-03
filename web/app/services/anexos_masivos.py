# web/app/services/anexos_masivos.py
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re

from docx import Document
from sqlalchemy.orm import joinedload
from flask import current_app

from ..config import generar_nombre_documento
from ..extensions import db
from ..models import Contrato, DocumentoLaboral, Obra, ParametroLaboral, Trabajador
from ..models.anexos_masivos import DetalleAnexoMasivo, ProcesoAnexoMasivo
from ..services.docx_engine import replace_placeholders
from ..services.paths_nextcloud import get_nc_paths
from ..services.pdf_convert import convert_docx_to_pdf
from ..services.storage_nextcloud_fs import ensure_dir
from ..utils import fecha_larga_es, format_rut_parts


def _safe_filename(name: str) -> str:
    name = (name or "").strip()
    name = re.sub(r'[\\\/:\*\?"<>\|]+', " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name.rstrip(" .")


def _norm_formato(value: str | None, default: str = "AMBOS") -> str:
    fmt = (value or default).strip().upper()
    if fmt not in ("DOCX", "PDF", "AMBOS"):
        return default
    return fmt


def _formato_clp_sin_decimales(valor) -> str:
    try:
        return f"{int(valor):,}".replace(",", ".")
    except Exception:
        return "0"


UNIDADES = (
    "", "uno", "dos", "tres", "cuatro", "cinco", "seis",
    "siete", "ocho", "nueve", "diez", "once", "doce",
    "trece", "catorce", "quince", "dieciséis", "diecisiete",
    "dieciocho", "diecinueve", "veinte", "veintiuno", "veintidós",
    "veintitrés", "veinticuatro", "veinticinco", "veintiséis",
    "veintisiete", "veintiocho", "veintinueve"
)

DECENAS = (
    "", "", "treinta", "cuarenta", "cincuenta",
    "sesenta", "setenta", "ochenta", "noventa"
)

CENTENAS = (
    "", "ciento", "doscientos", "trescientos", "cuatrocientos",
    "quinientos", "seiscientos", "setecientos", "ochocientos", "novecientos"
)


def numero_a_palabras_es(n: int) -> str:
    if n == 0:
        return "cero"
    if n == 100:
        return "cien"

    def _convert_triplet(num: int) -> str:
        c = num // 100
        d = (num % 100) // 10
        u = num % 10
        du = num % 100

        partes = []

        if c > 0:
            if num == 100:
                partes.append("cien")
            else:
                partes.append(CENTENAS[c])

        if du <= 29:
            if du > 0:
                partes.append(UNIDADES[du])
        else:
            texto_decena = DECENAS[d]
            if u > 0:
                texto_decena = f"{texto_decena} y {UNIDADES[u]}"
            partes.append(texto_decena)

        return " ".join(p for p in partes if p).strip()

    millones = n // 1_000_000
    miles = (n % 1_000_000) // 1_000
    cientos = n % 1_000

    partes = []

    if millones:
        if millones == 1:
            partes.append("un millón")
        else:
            partes.append(f"{numero_a_palabras_es(millones)} millones")

    if miles:
        if miles == 1:
            partes.append("mil")
        else:
            partes.append(f"{_convert_triplet(miles)} mil")

    if cientos:
        partes.append(_convert_triplet(cientos))

    return " ".join(partes).strip()


def obtener_parametro_vigente_para_fecha(fecha_referencia):
    return (
        ParametroLaboral.query
        .filter(
            ParametroLaboral.fecha_vigencia_desde <= fecha_referencia,
            ParametroLaboral.fecha_vigencia_hasta >= fecha_referencia,
        )
        .order_by(
            ParametroLaboral.fecha_vigencia_desde.desc(),
            ParametroLaboral.id.desc(),
        )
        .first()
    )


def obtener_candidatos_sueldo_minimo(*, sueldo_objetivo, horas_minimas=40, obra_ids=None):
    query = (
        Contrato.query
        .options(
            joinedload(Contrato.trabajador),
            joinedload(Contrato.empleador),
            joinedload(Contrato.obra),
            joinedload(Contrato.cargo),
            joinedload(Contrato.horario),
        )
        .join(Trabajador, Contrato.trabajador_id == Trabajador.id)
        .join(Obra, Contrato.obra_id == Obra.id)
        .filter(
            Trabajador.estado_trabajador == "VIGENTE",
            Contrato.estado_contrato == "VIGENTE",
            Contrato.horas_semanales.isnot(None),
            Contrato.horas_semanales >= horas_minimas,
            Contrato.sueldo_base.isnot(None),
            Contrato.sueldo_base < sueldo_objetivo,
        )
    )

    if obra_ids:
        query = query.filter(Contrato.obra_id.in_(obra_ids))

    return (
        query.order_by(
            Obra.nombre.asc(),
            Trabajador.ap_paterno.asc(),
            Trabajador.ap_materno.asc(),
            Trabajador.nombres.asc(),
        )
        .all()
    )


def crear_proceso_sueldo_minimo(
    *,
    nombre: str,
    fecha_anexo,
    fecha_vigencia,
    horas_minimas: int = 40,
    obra_ids=None,
    ley_s_minimo_nro: str | None = None,
    ley_s_minimo_fecha_publicacion=None,
    usar_articulo_tercero: bool = False,
    articulo_tercero_texto: str | None = None,
    formato_generacion: str = "AMBOS",
    descripcion: str | None = None,
    observaciones: str | None = None,
):
    parametro = obtener_parametro_vigente_para_fecha(fecha_vigencia)
    if not parametro:
        raise ValueError("No existe un parámetro laboral vigente para la fecha indicada.")

    sueldo_objetivo = parametro.renta_minima_dependiente
    if sueldo_objetivo is None:
        raise ValueError("El parámetro laboral vigente no tiene informada la renta mínima dependiente.")

    proceso = ProcesoAnexoMasivo(
        tipo_proceso="SUELDO_MINIMO",
        nombre=nombre,
        descripcion=descripcion,
        fecha_anexo=fecha_anexo,
        fecha_vigencia=fecha_vigencia,
        parametro_laboral_id=parametro.id,
        sueldo_base_objetivo=sueldo_objetivo,
        filtrar_solo_vigentes=True,
        horas_semanales_minimas=horas_minimas,
        solo_sueldos_inferiores_objetivo=True,
        ley_s_minimo_nro=(ley_s_minimo_nro or "").strip() or None,
        ley_s_minimo_fecha_publicacion=ley_s_minimo_fecha_publicacion,
        usar_articulo_tercero=bool(usar_articulo_tercero),
        articulo_tercero_texto=(articulo_tercero_texto or "").strip() or None,
        formato_generacion=_norm_formato(formato_generacion, default="AMBOS"),
        estado="BORRADOR",
        observaciones=observaciones,
    )
    db.session.add(proceso)
    db.session.flush()

    contratos = obtener_candidatos_sueldo_minimo(
        sueldo_objetivo=sueldo_objetivo,
        horas_minimas=horas_minimas,
        obra_ids=obra_ids,
    )

    for contrato in contratos:
        detalle = DetalleAnexoMasivo(
            proceso_id=proceso.id,
            trabajador_id=contrato.trabajador_id,
            contrato_id=contrato.id,
            horas_semanales=contrato.horas_semanales,
            sueldo_base_actual=contrato.sueldo_base,
            sueldo_base_nuevo=sueldo_objetivo,
            seleccionado=True,
            estado="PENDIENTE",
            formato=proceso.formato_generacion,
            meta={
                "generacion": {
                    "ultimo_resultado": None,
                    "ultimo_formato": None,
                    "ultimo_ts": None,
                }
            },
        )
        db.session.add(detalle)

    db.session.commit()
    return proceso


def marcar_detalles_proceso(proceso: ProcesoAnexoMasivo, detalle_ids_seleccionados: list[int]):
    seleccionados = set(detalle_ids_seleccionados or [])

    for detalle in proceso.detalles:
        detalle.seleccionado = detalle.id in seleccionados
        if not detalle.seleccionado and detalle.estado == "PENDIENTE":
            detalle.estado = "OMITIDO"
        elif detalle.seleccionado and detalle.estado == "OMITIDO":
            detalle.estado = "PENDIENTE"

    db.session.commit()


def _get_tpl_anexo_sueldo_minimo() -> Path:
    templates_root = Path(current_app.root_path) / "document_templates" / "anexos"
    return templates_root / "anexo_sueldo_minimo.docx"


def _build_variables_anexo_sueldo_minimo(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    contrato: Contrato,
) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    rut_digits = (getattr(trabajador, "rut", "") or "").strip()
    dv = (getattr(trabajador, "dv", "") or "").strip().upper()
    rut_trabajador = format_rut_parts(rut_digits, dv) or (f"{rut_digits}-{dv}" if rut_digits and dv else rut_digits)

    monto_nuevo_int = int(detalle.sueldo_base_nuevo or 0)

    art_tercero_label = "TERCERO:" if proceso.usar_articulo_tercero and proceso.articulo_tercero_texto else ""
    art_tercero_texto = (proceso.articulo_tercero_texto or "").strip() if art_tercero_label else ""

    return {
        "{{EMPLEADOR_RAZON_SOCIAL}}": empleador.razon_social if empleador else "",
        "{{Nombre Empleador}}": empleador.razon_social if empleador else "",
        "{{OBRA_NOMBRE}}": obra.nombre if obra else "",
        "{{Nombre Obra}}": obra.nombre if obra else "",
        "{{OBRA_COMUNA}}": obra.comuna if obra and obra.comuna else "",
        "{{ANEXO_S_MINIMO_FECHA}}": fecha_larga_es(proceso.fecha_anexo),
        "{{EMPLEADOR_RUT}}": getattr(empleador, "rut", "") or "",
        "{{REP_LEGAL_NOMBRE}}": getattr(empleador, "nombre_rep_legal", "") or "",
        "{{REP_LEGAL_RUT}}": getattr(empleador, "rut_rep_legal", "") or "",
        "{{EMPLEADOR_DIRECCION}}": getattr(empleador, "direccion", "") or "",
        "{{EMPLEADOR_COMUNA}}": getattr(empleador, "comuna", "") or "",
        "{{TRABAJADOR_NOMBRES}}": trabajador.nombres or "",
        "{{TRABAJADOR_AP_PATERNO}}": trabajador.ap_paterno or "",
        "{{TRABAJADOR_AP_MATERNO}}": trabajador.ap_materno or "",
        "{{TRABAJADOR_RUT}}": rut_trabajador,
        "{{LEY_S_MINIMO_NRO}}": proceso.ley_s_minimo_nro or "",
        "{{LEY_S_MINIMO_FECHA}}": fecha_larga_es(proceso.ley_s_minimo_fecha_publicacion) if proceso.ley_s_minimo_fecha_publicacion else "",
        "{{SUELDO_MINIMO_FECHA_VIGENCIA}}": fecha_larga_es(proceso.fecha_vigencia),
        "{{SUELDO_MINIMO_MONTO}}": _formato_clp_sin_decimales(detalle.sueldo_base_nuevo),
        "{{SUELDO_MINIMO_MONTO_PALABRAS}}": numero_a_palabras_es(monto_nuevo_int),
        "{{ART_TERCERO:}}": art_tercero_label,
        "{{ART_TERCERO_TEXTO}}": art_tercero_texto,
    }


def _build_docx(template_path: Path, variables: dict[str, str], output_path: Path) -> Path:
    doc = Document(str(template_path))
    replace_placeholders(doc, variables)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _nombre_anexo_sueldo_minimo(trabajador, fecha_ref: date, extension: str) -> str:
    ap_paterno = (getattr(trabajador, "ap_paterno", "") or "").strip() or "SIN_APELLIDO"
    ap_materno = (getattr(trabajador, "ap_materno", "") or "").strip()
    nombres = (getattr(trabajador, "nombres", "") or "").strip()

    nombre_persona = " ".join([x for x in [ap_paterno, ap_materno, nombres] if x]).strip()
    base = f"{fecha_ref.strftime('%Y-%m-%d')} Anexo sueldo mínimo - {nombre_persona}"
    return f"{_safe_filename(base)}.{(extension or '').lstrip('.').lower()}"


def _touch_meta_generacion(detalle: DetalleAnexoMasivo, *, formato: str, resultado: str, detalle_msg: str | None = None):
    meta = detalle.meta or {}
    if not isinstance(meta, dict):
        meta = {}

    gen = meta.get("generacion", {})
    if not isinstance(gen, dict):
        gen = {}

    gen["ultimo_formato"] = formato
    gen["ultimo_ts"] = datetime.now().isoformat(timespec="seconds")
    gen["ultimo_resultado"] = resultado
    if detalle_msg:
        gen["detalle"] = detalle_msg

    meta["generacion"] = gen
    detalle.meta = meta


def generar_archivos_detalle_sueldo_minimo(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    formato: str | None = None,
):
    contrato = detalle.contrato
    if not contrato:
        raise RuntimeError("Detalle sin contrato asociado.")
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")
    if not contrato.obra:
        raise RuntimeError("El contrato no tiene obra asignada.")
    if not getattr(contrato.obra, "centro_costo", None):
        raise RuntimeError(f"La obra '{contrato.obra.nombre}' no tiene centro de costo definido.")

    tpl = _get_tpl_anexo_sueldo_minimo()
    if not tpl.exists():
        raise RuntimeError(
            "No se encontró la plantilla 'anexo_sueldo_minimo.docx' en 'document_templates/anexos/'."
        )

    formato = _norm_formato(formato or detalle.formato or proceso.formato_generacion, default="AMBOS")

    trabajador = contrato.trabajador
    centro_costo = (contrato.obra.centro_costo or "").strip()

    variables = _build_variables_anexo_sueldo_minimo(
        proceso=proceso,
        detalle=detalle,
        contrato=contrato,
    )

    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_s_minimo_{proceso.id}_{detalle.id}_{stamp}.docx"
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)
    ensure_dir(dest_dir)

    _build_docx(tpl, variables, docx_tmp)

    fecha_ref = proceso.fecha_anexo or date.today()
    nombre_docx = _nombre_anexo_sueldo_minimo(trabajador=trabajador, fecha_ref=fecha_ref, extension="docx")
    nombre_pdf = _nombre_anexo_sueldo_minimo(trabajador=trabajador, fecha_ref=fecha_ref, extension="pdf")

    docx_final = dest_dir / nombre_docx
    pdf_final = dest_dir / nombre_pdf

    if formato in ("DOCX", "AMBOS"):
        docx_final.write_bytes(docx_tmp.read_bytes())

    if formato in ("PDF", "AMBOS"):
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())

    detalle.formato = formato
    if formato in ("DOCX", "AMBOS"):
        detalle.docx_nombre = nombre_docx
        detalle.docx_ruta = str(docx_final)

    if formato in ("PDF", "AMBOS"):
        detalle.pdf_nombre = nombre_pdf
        detalle.pdf_ruta = str(pdf_final)

    detalle.estado = "GENERADO"
    detalle.observacion = None

    _touch_meta_generacion(detalle, formato=formato, resultado="OK", detalle_msg=None)

    _registrar_documentos_laborales(detalle=detalle, contrato=contrato)


def _registrar_documentos_laborales(*, detalle: DetalleAnexoMasivo, contrato: Contrato):
    tipo_por_proceso = {
        "SUELDO_MINIMO": "ANEXO_SUELDO_MINIMO",
        "LEY_40_HORAS": "ANEXO_LEY_40_HORAS",
        "REEMPLAZO_DIA": "ANEXO_REEMPLAZO_DIA",
    }
    tipo_doc = tipo_por_proceso.get(detalle.proceso.tipo_proceso, "ANEXO_MASIVO")

    existentes = DocumentoLaboral.query.filter(
        DocumentoLaboral.contrato_id == contrato.id,
        DocumentoLaboral.tipo == tipo_doc,
    ).all()

    rutas_existentes = {d.ruta_archivo for d in existentes if d.ruta_archivo}

    nuevos = []

    if detalle.docx_ruta and detalle.docx_ruta not in rutas_existentes:
        docx_doc = DocumentoLaboral.crear_para_contrato(
            contrato=contrato,
            tipo=tipo_doc,
            fecha_ref=detalle.proceso.fecha_anexo,
            extension="docx",
            ruta_archivo=detalle.docx_ruta,
            estado="VIGENTE",
        )
        docx_doc.nombre_archivo = detalle.docx_nombre or docx_doc.nombre_archivo
        nuevos.append(docx_doc)

    if detalle.pdf_ruta and detalle.pdf_ruta not in rutas_existentes:
        pdf_doc = DocumentoLaboral.crear_para_contrato(
            contrato=contrato,
            tipo=tipo_doc,
            fecha_ref=detalle.proceso.fecha_anexo,
            extension="pdf",
            ruta_archivo=detalle.pdf_ruta,
            estado="VIGENTE",
        )
        pdf_doc.nombre_archivo = detalle.pdf_nombre or pdf_doc.nombre_archivo
        nuevos.append(pdf_doc)

    for doc in nuevos:
        db.session.add(doc)


def generar_documentos_proceso(proceso: ProcesoAnexoMasivo, formato: str | None = None):
    total_ok = 0
    total_error = 0

    for detalle in proceso.detalles:
        if not detalle.seleccionado:
            if detalle.estado in ("PENDIENTE", "OMITIDO"):
                detalle.estado = "OMITIDO"
                db.session.commit()
            continue

        try:
            _touch_meta_generacion(
                detalle,
                formato=_norm_formato(formato or detalle.formato or proceso.formato_generacion),
                resultado="ERROR",
                detalle_msg="Inicio generación (pendiente)",
            )
            db.session.commit()

            if proceso.tipo_proceso == "LEY_40_HORAS":
                generar_archivos_detalle_ley_40_horas(
                    proceso=proceso,
                    detalle=detalle,
                    formato=formato,
                )
            elif proceso.tipo_proceso == "REEMPLAZO_DIA":
                generar_archivos_detalle_reemplazo_dia(
                    proceso=proceso,
                    detalle=detalle,
                    formato=formato,
                )
            else:
                generar_archivos_detalle_sueldo_minimo(
                    proceso=proceso,
                    detalle=detalle,
                    formato=formato,
                )

            db.session.commit()
            total_ok += 1

        except Exception as e:
            db.session.rollback()

            try:
                detalle = DetalleAnexoMasivo.query.get(detalle.id)
                if detalle:
                    detalle.estado = "ERROR"
                    detalle.observacion = str(e)
                    _touch_meta_generacion(
                        detalle,
                        formato=_norm_formato(formato or detalle.formato or proceso.formato_generacion),
                        resultado="ERROR",
                        detalle_msg=str(e),
                    )
                    db.session.commit()
            except Exception:
                db.session.rollback()

            total_error += 1

    proceso = ProcesoAnexoMasivo.query.get(proceso.id)
    if proceso:
        if total_ok > 0 and total_error == 0:
            proceso.estado = "GENERADO"
        elif total_ok > 0 and total_error > 0:
            proceso.estado = "BORRADOR"
        db.session.commit()

    return total_ok, total_error


def aplicar_cambio_sueldos_proceso(proceso: ProcesoAnexoMasivo):
    cambios = 0

    for detalle in proceso.detalles:
        if not detalle.seleccionado:
            if detalle.estado in ("PENDIENTE", "OMITIDO"):
                detalle.estado = "OMITIDO"
            continue

        contrato = Contrato.query.get(detalle.contrato_id)
        if not contrato:
            detalle.estado = "ERROR"
            detalle.observacion = "Contrato no encontrado al aplicar cambio."
            continue

        contrato.sueldo_base = detalle.sueldo_base_nuevo
        detalle.estado = "APLICADO"
        cambios += 1

    proceso.estado = "APLICADO"
    db.session.commit()
    return cambios

def obtener_candidatos_ley_40_horas(*, horas_actuales_minimas=44, obra_ids=None):
    query = (
        Contrato.query
        .options(
            joinedload(Contrato.trabajador),
            joinedload(Contrato.empleador),
            joinedload(Contrato.obra),
            joinedload(Contrato.cargo),
            joinedload(Contrato.horario),
        )
        .join(Trabajador, Contrato.trabajador_id == Trabajador.id)
        .join(Obra, Contrato.obra_id == Obra.id)
        .filter(
            Trabajador.estado_trabajador == "VIGENTE",
            Contrato.estado_contrato == "VIGENTE",
            Contrato.horas_semanales.isnot(None),
            Contrato.horas_semanales >= horas_actuales_minimas,
        )
    )

    if obra_ids:
        query = query.filter(Contrato.obra_id.in_(obra_ids))

    return (
        query.order_by(
            Obra.nombre.asc(),
            Trabajador.ap_paterno.asc(),
            Trabajador.ap_materno.asc(),
            Trabajador.nombres.asc(),
        )
        .all()
    )

def crear_proceso_ley_40_horas(
    *,
    nombre: str,
    fecha_anexo,
    fecha_vigencia,
    horas_actuales_minimas: int = 44,
    horas_nuevas: int = 42,
    obra_ids=None,
    ley_40_hrs_nro: str | None = None,
    ley_40_hrs_fecha=None,
    formato_generacion: str = "AMBOS",
    descripcion: str | None = None,
    observaciones: str | None = None,
):
    proceso = ProcesoAnexoMasivo(
        tipo_proceso="LEY_40_HORAS",
        nombre=nombre,
        descripcion=descripcion,
        fecha_anexo=fecha_anexo,
        fecha_vigencia=fecha_vigencia,

        # Estos campos quedan neutros para cumplir la estructura actual de la tabla
        parametro_laboral_id=None,
        sueldo_base_objetivo=0,
        filtrar_solo_vigentes=True,
        horas_semanales_minimas=horas_actuales_minimas,
        solo_sueldos_inferiores_objetivo=False,

        # Reutilizamos estos campos existentes para la ley 40 horas
        ley_s_minimo_nro=(ley_40_hrs_nro or "").strip() or None,
        ley_s_minimo_fecha_publicacion=ley_40_hrs_fecha,

        usar_articulo_tercero=False,
        articulo_tercero_texto=None,
        formato_generacion=_norm_formato(formato_generacion, default="AMBOS"),
        estado="BORRADOR",
        observaciones=observaciones,
    )
    db.session.add(proceso)
    db.session.flush()

    contratos = obtener_candidatos_ley_40_horas(
        horas_actuales_minimas=horas_actuales_minimas,
        obra_ids=obra_ids,
    )

    for contrato in contratos:
        sueldo_actual = contrato.sueldo_base or 0

        detalle = DetalleAnexoMasivo(
            proceso_id=proceso.id,
            trabajador_id=contrato.trabajador_id,
            contrato_id=contrato.id,
            horas_semanales=contrato.horas_semanales,
            sueldo_base_actual=sueldo_actual,
            sueldo_base_nuevo=sueldo_actual,
            seleccionado=True,
            estado="PENDIENTE",
            formato=proceso.formato_generacion,
            meta={
                "horas_nuevas": horas_nuevas,
                "generacion": {
                    "ultimo_resultado": None,
                    "ultimo_formato": None,
                    "ultimo_ts": None,
                }
            },
        )
        db.session.add(detalle)

    db.session.commit()
    return proceso

def _get_tpl_anexo_ley_40_horas() -> Path:
    templates_root = Path(current_app.root_path) / "document_templates" / "anexos"
    return templates_root / "anexo_42_hrs.docx"

def _build_variables_anexo_ley_40_horas(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    contrato: Contrato,
) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    rut_digits = (getattr(trabajador, "rut", "") or "").strip()
    dv = (getattr(trabajador, "dv", "") or "").strip().upper()
    rut_trabajador = format_rut_parts(rut_digits, dv) or (f"{rut_digits}-{dv}" if rut_digits and dv else rut_digits)

    return {
        "{{EMPLEADOR_RAZON_SOCIAL}}": empleador.razon_social if empleador else "",
        "{{Nombre Empleador}}": empleador.razon_social if empleador else "",
        "{{OBRA_NOMBRE}}": obra.nombre if obra else "",
        "{{Nombre Obra}}": obra.nombre if obra else "",
        "{{OBRA_COMUNA}}": obra.comuna if obra and obra.comuna else "",
        "{{ANEXO_LEY_40_HRS_FECHA}}": fecha_larga_es(proceso.fecha_anexo),
        "{{EMPLEADOR_RUT}}": getattr(empleador, "rut", "") or "",
        "{{REP_LEGAL_NOMBRE}}": getattr(empleador, "nombre_rep_legal", "") or "",
        "{{REP_LEGAL_RUT}}": getattr(empleador, "rut_rep_legal", "") or "",
        "{{EMPLEADOR_DIRECCION}}": getattr(empleador, "direccion", "") or "",
        "{{EMPLEADOR_COMUNA}}": getattr(empleador, "comuna", "") or "",
        "{{TRABAJADOR_NOMBRES}}": trabajador.nombres or "",
        "{{TRABAJADOR_AP_PATERNO}}": trabajador.ap_paterno or "",
        "{{TRABAJADOR_AP_MATERNO}}": trabajador.ap_materno or "",
        "{{TRABAJADOR_RUT}}": rut_trabajador,
        "{{LEY_40_HRS_NRO}}": proceso.ley_s_minimo_nro or "",
        "{{LEY_40_HRS_FECHA}}": fecha_larga_es(proceso.ley_s_minimo_fecha_publicacion) if proceso.ley_s_minimo_fecha_publicacion else "",
    }

def _nombre_anexo_ley_40_horas(trabajador, fecha_ref: date, extension: str) -> str:
    ap_paterno = (getattr(trabajador, "ap_paterno", "") or "").strip() or "SIN_APELLIDO"
    ap_materno = (getattr(trabajador, "ap_materno", "") or "").strip()
    nombres = (getattr(trabajador, "nombres", "") or "").strip()

    nombre_persona = " ".join([x for x in [ap_paterno, ap_materno, nombres] if x]).strip()
    base = f"{fecha_ref.strftime('%Y-%m-%d')} Anexo ley 40 horas - {nombre_persona}"
    return f"{_safe_filename(base)}.{(extension or '').lstrip('.').lower()}"

def generar_archivos_detalle_ley_40_horas(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    formato: str | None = None,
):
    contrato = detalle.contrato
    if not contrato:
        raise RuntimeError("Detalle sin contrato asociado.")
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")
    if not contrato.obra:
        raise RuntimeError("El contrato no tiene obra asignada.")
    if not getattr(contrato.obra, "centro_costo", None):
        raise RuntimeError(f"La obra '{contrato.obra.nombre}' no tiene centro de costo definido.")

    tpl = _get_tpl_anexo_ley_40_horas()
    if not tpl.exists():
        raise RuntimeError(
            "No se encontró la plantilla 'anexo_42_hrs.docx' en 'document_templates/anexos/'."
        )

    formato = _norm_formato(formato or detalle.formato or proceso.formato_generacion, default="AMBOS")

    trabajador = contrato.trabajador
    centro_costo = (contrato.obra.centro_costo or "").strip()

    variables = _build_variables_anexo_ley_40_horas(
        proceso=proceso,
        detalle=detalle,
        contrato=contrato,
    )

    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_42hrs_{proceso.id}_{detalle.id}_{stamp}.docx"
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)
    ensure_dir(dest_dir)

    _build_docx(tpl, variables, docx_tmp)

    fecha_ref = proceso.fecha_anexo or date.today()
    nombre_docx = _nombre_anexo_ley_40_horas(trabajador=trabajador, fecha_ref=fecha_ref, extension="docx")
    nombre_pdf = _nombre_anexo_ley_40_horas(trabajador=trabajador, fecha_ref=fecha_ref, extension="pdf")

    docx_final = dest_dir / nombre_docx
    pdf_final = dest_dir / nombre_pdf

    if formato in ("DOCX", "AMBOS"):
        docx_final.write_bytes(docx_tmp.read_bytes())

    if formato in ("PDF", "AMBOS"):
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())

    detalle.formato = formato
    if formato in ("DOCX", "AMBOS"):
        detalle.docx_nombre = nombre_docx
        detalle.docx_ruta = str(docx_final)

    if formato in ("PDF", "AMBOS"):
        detalle.pdf_nombre = nombre_pdf
        detalle.pdf_ruta = str(pdf_final)

    detalle.estado = "GENERADO"
    detalle.observacion = None

    _touch_meta_generacion(detalle, formato=formato, resultado="OK", detalle_msg=None)

    _registrar_documentos_laborales(detalle=detalle, contrato=contrato)


def obtener_candidatos_reemplazo_dia(*, obra_ids=None):
    """Contratos vigentes disponibles para documentar acuerdo de reemplazo de día."""
    query = (
        Contrato.query
        .options(
            joinedload(Contrato.trabajador),
            joinedload(Contrato.empleador),
            joinedload(Contrato.obra),
            joinedload(Contrato.cargo),
            joinedload(Contrato.horario),
        )
        .join(Trabajador, Contrato.trabajador_id == Trabajador.id)
        .join(Obra, Contrato.obra_id == Obra.id)
        .filter(
            Trabajador.estado_trabajador == "VIGENTE",
            Contrato.estado_contrato == "VIGENTE",
        )
    )

    if obra_ids:
        query = query.filter(Contrato.obra_id.in_(obra_ids))

    return (
        query.order_by(
            Obra.nombre.asc(),
            Trabajador.ap_paterno.asc(),
            Trabajador.ap_materno.asc(),
            Trabajador.nombres.asc(),
        )
        .all()
    )


def crear_proceso_reemplazo_dia(
    *,
    nombre: str,
    fecha_anexo,
    fecha_trabajar,
    fecha_reemplazar,
    obra_ids=None,
    formato_generacion: str = "AMBOS",
    descripcion: str | None = None,
    observaciones: str | None = None,
):
    if fecha_trabajar == fecha_reemplazar:
        raise ValueError("El día a trabajar y el día a reemplazar no pueden ser el mismo.")

    proceso = ProcesoAnexoMasivo(
        tipo_proceso="REEMPLAZO_DIA",
        nombre=nombre,
        descripcion=descripcion,
        fecha_anexo=fecha_anexo,
        # Reutilizamos fecha_vigencia como día que se trabajará excepcionalmente.
        fecha_vigencia=fecha_trabajar,
        parametro_laboral_id=None,
        sueldo_base_objetivo=0,
        filtrar_solo_vigentes=True,
        horas_semanales_minimas=0,
        solo_sueldos_inferiores_objetivo=False,
        ley_s_minimo_nro=None,
        # Reutilizamos este campo como día interferiado/reemplazado para evitar migración.
        ley_s_minimo_fecha_publicacion=fecha_reemplazar,
        usar_articulo_tercero=False,
        articulo_tercero_texto=None,
        formato_generacion=_norm_formato(formato_generacion, default="AMBOS"),
        estado="BORRADOR",
        observaciones=observaciones,
    )
    db.session.add(proceso)
    db.session.flush()

    contratos = obtener_candidatos_reemplazo_dia(obra_ids=obra_ids)

    for contrato in contratos:
        sueldo_actual = contrato.sueldo_base or 0
        detalle = DetalleAnexoMasivo(
            proceso_id=proceso.id,
            trabajador_id=contrato.trabajador_id,
            contrato_id=contrato.id,
            horas_semanales=contrato.horas_semanales,
            sueldo_base_actual=sueldo_actual,
            sueldo_base_nuevo=sueldo_actual,
            seleccionado=True,
            estado="PENDIENTE",
            formato=proceso.formato_generacion,
            meta={
                "fecha_trabajar": fecha_trabajar.isoformat() if fecha_trabajar else None,
                "fecha_reemplazar": fecha_reemplazar.isoformat() if fecha_reemplazar else None,
                "generacion": {
                    "ultimo_resultado": None,
                    "ultimo_formato": None,
                    "ultimo_ts": None,
                },
            },
        )
        db.session.add(detalle)

    db.session.commit()
    return proceso


def _get_tpl_anexo_reemplazo_dia() -> Path:
    templates_root = Path(current_app.root_path) / "document_templates" / "anexos"
    return templates_root / "anexo_reemplazo_dia.docx"


def _build_variables_anexo_reemplazo_dia(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    contrato: Contrato,
) -> dict[str, str]:
    trabajador = contrato.trabajador
    empleador = contrato.empleador
    obra = contrato.obra

    rut_digits = (getattr(trabajador, "rut", "") or "").strip()
    dv = (getattr(trabajador, "dv", "") or "").strip().upper()
    rut_trabajador = format_rut_parts(rut_digits, dv) or (f"{rut_digits}-{dv}" if rut_digits and dv else rut_digits)

    fecha_trabajar = proceso.fecha_vigencia
    fecha_reemplazar = proceso.ley_s_minimo_fecha_publicacion

    return {
        "{{EMPLEADOR_RAZON_SOCIAL}}": empleador.razon_social if empleador else "",
        "{{Nombre Empleador}}": empleador.razon_social if empleador else "",
        "{{OBRA_NOMBRE}}": obra.nombre if obra else "",
        "{{Nombre Obra}}": obra.nombre if obra else "",
        "{{OBRA_COMUNA}}": obra.comuna if obra and obra.comuna else "",
        "{{ANEXO_CAMBIO_DIA_FECHA}}": fecha_larga_es(proceso.fecha_anexo),
        "{{EMPLEADOR_RUT}}": getattr(empleador, "rut", "") or "",
        "{{REP_LEGAL_NOMBRE}}": getattr(empleador, "nombre_rep_legal", "") or "",
        "{{REP_LEGAL_RUT}}": getattr(empleador, "rut_rep_legal", "") or "",
        "{{EMPLEADOR_DIRECCION}}": getattr(empleador, "direccion", "") or "",
        "{{EMPLEADOR_COMUNA}}": getattr(empleador, "comuna", "") or "",
        "{{TRABAJADOR_NOMBRES}}": trabajador.nombres or "",
        "{{TRABAJADOR_AP_PATERNO}}": trabajador.ap_paterno or "",
        "{{TRABAJADOR_AP_MATERNO}}": trabajador.ap_materno or "",
        "{{TRABAJADOR_RUT}}": rut_trabajador,
        "{{ANEXO_CAMBIO_DIA_TRABAJAR}}": fecha_larga_es(fecha_trabajar),
        "{{ANEXO_CAMBIO_DIA_REEMPLAZAR}}": fecha_larga_es(fecha_reemplazar),
    }


def _nombre_anexo_reemplazo_dia(trabajador, fecha_ref: date, fecha_trabajar: date | None, extension: str) -> str:
    ap_paterno = (getattr(trabajador, "ap_paterno", "") or "").strip() or "SIN_APELLIDO"
    ap_materno = (getattr(trabajador, "ap_materno", "") or "").strip()
    nombres = (getattr(trabajador, "nombres", "") or "").strip()

    nombre_persona = " ".join([x for x in [ap_paterno, ap_materno, nombres] if x]).strip()
    fecha_trabajo_txt = fecha_trabajar.strftime("%Y-%m-%d") if fecha_trabajar else fecha_ref.strftime("%Y-%m-%d")
    base = f"{fecha_ref.strftime('%Y-%m-%d')} Anexo reemplazo día {fecha_trabajo_txt} - {nombre_persona}"
    return f"{_safe_filename(base)}.{(extension or '').lstrip('.').lower()}"


def generar_archivos_detalle_reemplazo_dia(
    *,
    proceso: ProcesoAnexoMasivo,
    detalle: DetalleAnexoMasivo,
    formato: str | None = None,
):
    contrato = detalle.contrato
    if not contrato:
        raise RuntimeError("Detalle sin contrato asociado.")
    if not contrato.trabajador:
        raise RuntimeError("Contrato sin trabajador asociado.")
    if not contrato.empleador:
        raise RuntimeError("El contrato no tiene empleador asignado.")
    if not contrato.obra:
        raise RuntimeError("El contrato no tiene obra asignada.")
    if not getattr(contrato.obra, "centro_costo", None):
        raise RuntimeError(f"La obra '{contrato.obra.nombre}' no tiene centro de costo definido.")

    tpl = _get_tpl_anexo_reemplazo_dia()
    if not tpl.exists():
        raise RuntimeError(
            "No se encontró la plantilla 'anexo_reemplazo_dia.docx' en 'document_templates/anexos/'."
        )

    formato = _norm_formato(formato or detalle.formato or proceso.formato_generacion, default="AMBOS")

    trabajador = contrato.trabajador
    centro_costo = (contrato.obra.centro_costo or "").strip()

    variables = _build_variables_anexo_reemplazo_dia(
        proceso=proceso,
        detalle=detalle,
        contrato=contrato,
    )

    nc = get_nc_paths()
    dest_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=trabajador.nombres or "",
        ap_paterno=trabajador.ap_paterno or "",
        ap_materno=trabajador.ap_materno or "",
        tipo_doc="ANEXOS",
    )

    out_tmp = Path(current_app.instance_path) / "generated_docs"
    ensure_dir(out_tmp)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    docx_tmp = out_tmp / f"tmp_anexo_reemplazo_dia_{proceso.id}_{detalle.id}_{stamp}.docx"
    pdf_tmp_dir = out_tmp / "pdf"
    ensure_dir(pdf_tmp_dir)
    ensure_dir(dest_dir)

    _build_docx(tpl, variables, docx_tmp)

    fecha_ref = proceso.fecha_anexo or date.today()
    fecha_trabajar = proceso.fecha_vigencia
    nombre_docx = _nombre_anexo_reemplazo_dia(
        trabajador=trabajador,
        fecha_ref=fecha_ref,
        fecha_trabajar=fecha_trabajar,
        extension="docx",
    )
    nombre_pdf = _nombre_anexo_reemplazo_dia(
        trabajador=trabajador,
        fecha_ref=fecha_ref,
        fecha_trabajar=fecha_trabajar,
        extension="pdf",
    )

    docx_final = dest_dir / nombre_docx
    pdf_final = dest_dir / nombre_pdf

    if formato in ("DOCX", "AMBOS"):
        docx_final.write_bytes(docx_tmp.read_bytes())

    if formato in ("PDF", "AMBOS"):
        pdf_tmp = convert_docx_to_pdf(docx_tmp, pdf_tmp_dir)
        pdf_final.write_bytes(pdf_tmp.read_bytes())

    detalle.formato = formato
    if formato in ("DOCX", "AMBOS"):
        detalle.docx_nombre = nombre_docx
        detalle.docx_ruta = str(docx_final)

    if formato in ("PDF", "AMBOS"):
        detalle.pdf_nombre = nombre_pdf
        detalle.pdf_ruta = str(pdf_final)

    detalle.estado = "GENERADO"
    detalle.observacion = None

    _touch_meta_generacion(detalle, formato=formato, resultado="OK", detalle_msg=None)
    _registrar_documentos_laborales(detalle=detalle, contrato=contrato)


def aplicar_cambios_proceso(proceso: ProcesoAnexoMasivo):
    if proceso.tipo_proceso == "LEY_40_HORAS":
        return aplicar_cambio_ley_40_horas_proceso(proceso)
    if proceso.tipo_proceso == "REEMPLAZO_DIA":
        return aplicar_cambio_reemplazo_dia_proceso(proceso)
    return aplicar_cambio_sueldos_proceso(proceso)


def aplicar_cambio_reemplazo_dia_proceso(proceso: ProcesoAnexoMasivo):
    """
    Este proceso es documental: no modifica campos del contrato.
    Al aplicar, solo marca los detalles seleccionados como APLICADO para cerrar el flujo.
    """
    cambios = 0

    for detalle in proceso.detalles:
        if not detalle.seleccionado:
            if detalle.estado in ("PENDIENTE", "OMITIDO"):
                detalle.estado = "OMITIDO"
            continue

        detalle.estado = "APLICADO"
        cambios += 1

    proceso.estado = "APLICADO"
    db.session.commit()
    return cambios


def aplicar_cambio_ley_40_horas_proceso(proceso: ProcesoAnexoMasivo):
    cambios = 0

    for detalle in proceso.detalles:
        if not detalle.seleccionado:
            if detalle.estado in ("PENDIENTE", "OMITIDO"):
                detalle.estado = "OMITIDO"
            continue

        contrato = Contrato.query.get(detalle.contrato_id)
        if not contrato:
            detalle.estado = "ERROR"
            detalle.observacion = "Contrato no encontrado al aplicar cambio."
            continue

        horas_nuevas = None
        if isinstance(detalle.meta, dict):
            horas_nuevas = detalle.meta.get("horas_nuevas")

        if not horas_nuevas:
            detalle.estado = "ERROR"
            detalle.observacion = "No se encontró la nueva jornada en el detalle."
            continue

        contrato.horas_semanales = horas_nuevas
        detalle.estado = "APLICADO"
        cambios += 1

    proceso.estado = "APLICADO"
    db.session.commit()
    return cambios
