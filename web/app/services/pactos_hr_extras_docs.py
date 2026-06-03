# web/app/services/pactos_hr_extras_docs.py
from __future__ import annotations

import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

from flask import current_app
from docx import Document as DocxDocument

from ..models import PactoHorasExtra
from .docx_engine import replace_placeholders
from .paths_nextcloud import get_nc_paths
from .rut_format import rut_con_puntos


TEMPLATE_REL_PATH = "document_templates/extras/pacto_hr_ex.docx"


def _fmt_fecha_es(d: date | None) -> str:
    if not d:
        return ""
    meses = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
    ]
    return f"{d.day} de {meses[d.month - 1]} de {d.year}"


def _safe_filename(s: str) -> str:
    keep = []
    for ch in (s or ""):
        if ch.isalnum() or ch in (" ", "-", "_", "."):
            keep.append(ch)
    out = "".join(keep).strip()
    out = re.sub(r"\s+", " ", out)
    return out or "PACTO_HORAS_EXTRA"


def _docx_to_pdf(docx_path: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)

    cmd = [
        "soffice",
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--norestore",
        "--convert-to", "pdf",
        "--outdir", str(out_dir),
        str(docx_path),
    ]

    try:
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    except FileNotFoundError as e:
        raise RuntimeError("No se encontró 'soffice'. Instala LibreOffice en el contenedor web.") from e
    except subprocess.CalledProcessError as e:
        err = (e.stderr or b"").decode("utf-8", errors="ignore")
        raise RuntimeError(f"Error convirtiendo DOCX a PDF (soffice): {err}") from e

    pdf_path = out_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError("La conversión a PDF terminó sin error, pero no se encontró el PDF generado.")
    return pdf_path


def _guardar_en_carpeta_trabajador_pacto(
    pacto: PactoHorasExtra,
    docx_path: Path,
    pdf_path: Path | None,
) -> tuple[str | None, str | None]:
    """
    Copia el DOCX/PDF a:
      /<CC>/Laboral/Trabajadores/<Ap Pat> <Ap Mat> <Nombres>/PACTOS_HORAS_EXTRA/
    """
    if pacto.trabajador is None:
        return (None, None)

    t = pacto.trabajador

    centro_costo = ""
    if pacto.obra is not None:
        centro_costo = (getattr(pacto.obra, "centro_costo", "") or "").strip()

    if not centro_costo:
        centro_costo = "SIN_NOMBRE"

    nc = get_nc_paths()
    destino_dir = nc.dir_trabajador(
        centro_costo=centro_costo,
        nombres=(t.nombres or ""),
        ap_paterno=(t.ap_paterno or ""),
        ap_materno=(t.ap_materno or ""),
        tipo_doc="PACTOS_HORAS_EXTRA",
    )

    destino_dir.mkdir(parents=True, exist_ok=True)

    dst_docx = destino_dir / docx_path.name
    shutil.copy2(docx_path, dst_docx)

    dst_pdf = None
    if pdf_path is not None:
        dst_pdf = destino_dir / pdf_path.name
        shutil.copy2(pdf_path, dst_pdf)

    return (str(dst_docx), str(dst_pdf) if dst_pdf else None)


def generar_docx_y_pdf_pacto_horas_extra(
    pacto: PactoHorasExtra,
) -> tuple[str, str, str | None, str | None, str | None, str | None]:
    """
    Genera DOCX desde plantilla, genera PDF y copia a Nextcloud.
    Retorna:
      (docx_tmp_path, docx_nombre, pdf_tmp_path, pdf_nombre, docx_guardado_path, pdf_guardado_path)
    """
    app_root = Path(current_app.root_path)  # web/app
    tpl_path = app_root / TEMPLATE_REL_PATH
    if not tpl_path.exists():
        raise FileNotFoundError(f"No existe plantilla: {tpl_path}")

    doc = DocxDocument(str(tpl_path))

    t = pacto.trabajador
    c = pacto.contrato
    emp = pacto.empleador
    obra = pacto.obra

    if t is None:
        raise RuntimeError("El pacto no tiene trabajador asociado.")
    if c is None:
        raise RuntimeError("El pacto no tiene contrato asociado.")

    emp_razon_social = (getattr(emp, "razon_social", "") or "") if emp else ""
    emp_rut = rut_con_puntos((getattr(emp, "rut", "") or "").strip()) if emp else ""
    rep_legal_nombre = (getattr(emp, "nombre_rep_legal", "") or "") if emp else ""
    rep_legal_rut = rut_con_puntos((getattr(emp, "rut_rep_legal", "") or "").strip()) if emp else ""
    emp_direccion = (getattr(emp, "direccion", "") or "") if emp else ""
    emp_comuna = (getattr(emp, "comuna", "") or "") if emp else ""

    obra_nombre = (getattr(obra, "nombre", "") or "") if obra else ""
    obra_comuna = (getattr(obra, "comuna", "") or "") if obra else ""

    trab_nombres = (t.nombres or "").strip()
    trab_ap_pat = (t.ap_paterno or "").strip()
    trab_ap_mat = (t.ap_materno or "").strip()
    trab_rut = rut_con_puntos(getattr(t, "rut_completo", None) or getattr(t, "rut", None) or "")

    placeholders = {
        "EMPLEADOR_RAZON_SOCIAL": emp_razon_social,
        "EMPLEADOR_RUT": emp_rut,
        "REP_LEGAL_NOMBRE": rep_legal_nombre,
        "REP_LEGAL_RUT": rep_legal_rut,
        "EMPLEADOR_DIRECCION": emp_direccion,
        "EMPLEADOR_COMUNA": emp_comuna,

        "OBRA_NOMBRE": obra_nombre,
        "OBRA_COMUNA": obra_comuna,

        "TRABAJADOR_NOMBRES": trab_nombres,
        "TRABAJADOR_AP_PATERNO": trab_ap_pat,
        "TRABAJADOR_AP_MATERNO": trab_ap_mat,
        "TRABAJADOR_RUT": trab_rut,

        "PACTO_FECHA_INICIO": _fmt_fecha_es(pacto.fecha_inicio),
        "PACTO_FECHA_TERMINO": _fmt_fecha_es(pacto.fecha_termino),
    }

    replace_placeholders(doc, placeholders)

    # Naming corporativo
    fecha_iso = (pacto.fecha_inicio or date.today()).isoformat()
    full_name = " ".join(
        p for p in [trab_ap_pat, trab_ap_mat, trab_nombres] if p
    ).strip()
    full_name = _safe_filename(full_name)

    docx_name = re.sub(
        r"\s+",
        " ",
        f"{fecha_iso} Pacto Horas Extra - {full_name}.docx"
    ).strip()

    pdf_name = re.sub(
        r"\s+",
        " ",
        f"{fecha_iso} Pacto Horas Extra - {full_name}.pdf"
    ).strip()

    # Guardado local temporal
    out_dir = Path("/tmp/rrhh_docs/pactos_horas_extra")
    out_dir.mkdir(parents=True, exist_ok=True)

    docx_path = out_dir / docx_name
    doc.save(str(docx_path))

    # PDF
    pdf_tmp_path = _docx_to_pdf(docx_path, out_dir)

    pdf_path = out_dir / pdf_name
    if pdf_tmp_path != pdf_path:
        try:
            if pdf_path.exists():
                pdf_path.unlink()
            pdf_tmp_path.rename(pdf_path)
        except Exception:
            shutil.copy2(pdf_tmp_path, pdf_path)

    # Copia a Nextcloud
    docx_guardado, pdf_guardado = _guardar_en_carpeta_trabajador_pacto(
        pacto=pacto,
        docx_path=docx_path,
        pdf_path=pdf_path,
    )

    return (
        str(docx_path), docx_name,
        str(pdf_path), pdf_name,
        docx_guardado, pdf_guardado,
    )